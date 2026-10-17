#!/usr/bin/env python3
"""Gera os 4 formulários simplificados de estudo (uma aba, 4 colunas).

    python3 -I scripts/formularios/gerar_formularios.py [pasta-de-saida]

Saída padrão: referencia/formularios/ — um .xlsx por tipo de estudo:
preliminar-incorporacao, preliminar-loteamento, avancado-incorporacao,
avancado-loteamento.

Cada arquivo tem UMA aba, "Formulário", com quatro colunas:
    A Campo · B Valor · C Unidade · D Chave técnica
e duas zonas:
    · PREENCHER (células amarelas, coluna B) — todas as entradas do estudo;
    · REFERÊNCIA CALCULADA (células cinzas) — fórmulas do Excel, só para conferência.

Regras de projeto (decisões do pedido que originou este gerador):
    · Campo que o app aceita em várias unidades (R$/m², R$ total, % do VGV…) vira UMA LINHA
      POR UNIDADE, todas em branco e todas manuais. O Excel nunca calcula uma unidade a partir
      de outra dentro da zona de entrada; os equivalentes aparecem só na zona calculada.
    · Percentual é percentual de verdade do Excel (4% = 0,04, formatado 0,00%). Digite sempre com o
      símbolo % (4%); o número solto fica ambíguo conforme o programa que abre a planilha.
    · Listas (produtos, tipologias, custos…) são achatadas em linhas numeradas "Item N — campo".
"""
import re
import sys
import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

# ── formatos ────────────────────────────────────────────────────────────────
F_TXT = '@'
F_M2 = '#,##0.00'
F_RS = '"R$" #,##0.00'
F_PCT = '0.00%'
F_X = '0.00'
F_INT = '0'
F_DATA = 'mmm/yyyy'
F_GERAL = 'General'

# ── paleta (arquivo de impressão/planilha: cores literais são o normal aqui) ─
COR_TITULO = '1F3864'
COR_SECAO = '2F5597'
COR_GRUPO = 'D9E2F3'
COR_ENTRADA = 'FFF2CC'
COR_FIXO = 'E7E6E6'
COR_CALC = 'EDEDED'
COR_CALC_SECAO = '595959'
COR_ALERTA = 'C00000'

LADO = Side(style='thin', color='BFBFBF')
BORDA = Border(left=LADO, right=LADO, top=LADO, bottom=LADO)


class Form:
    def __init__(self, titulo, subtitulo):
        self.titulo = titulo
        self.subtitulo = subtitulo
        self.linhas = []          # dicts
        self.nomes = {}           # nome -> índice em self.linhas
        self.exemplo = {}         # nome -> valor pré-preenchido (nos modelos em branco, só os padrões declarados em `ex=` de campos de escolha fixa)
        self.listas = {}          # nome da lista -> itens (vão para a coluna oculta F)

    # ── construtores de linha ──
    def secao(self, texto, calc=False):
        self.linhas.append(dict(t='secao', texto=texto, calc=calc))

    def grupo(self, texto):
        self.linhas.append(dict(t='grupo', texto=texto))

    def nota(self, texto):
        self.linhas.append(dict(t='nota', texto=texto))

    def _reg(self, nome, linha):
        if nome in self.nomes:
            raise ValueError(f'nome duplicado: {nome}')
        self.nomes[nome] = len(self.linhas)
        self.linhas.append(linha)

    def campo(self, nome, label, unidade, chave, fmt=F_GERAL, *, lista=None, minimo=None,
              maximo=None, fixo=None, maxlen=None, inteiro=False, ex=None, lista_ref=None,
              aviso=False):
        """lista = itens inline (<=255 caracteres); lista_ref = nome de lista auxiliar
        (coluna oculta); aviso=True deixa a validação como aviso (aceita texto fora da lista)."""
        self._reg(nome, dict(t='campo', label=label, unidade=unidade, chave=chave, fmt=fmt,
                             lista=lista, minimo=minimo, maximo=maximo, fixo=fixo,
                             maxlen=maxlen, inteiro=inteiro, lista_ref=lista_ref, aviso=aviso))
        if ex is not None:
            self.exemplo[nome] = ex

    def calc(self, nome, label, unidade, formula, fmt=F_GERAL, chave='calculado'):
        self._reg(nome, dict(t='calc', label=label, unidade=unidade, chave=chave,
                             fmt=fmt, formula=formula))


# ── atalhos de fórmula ──────────────────────────────────────────────────────
def r(nome):
    return '{' + nome + '}'


def soma(*nomes):
    return '+'.join(r(n) for n in nomes) if nomes else '0'


def vazio(nome):
    return f'{r(nome)}=""'


def cheio(nome):
    return f'{r(nome)}<>""'


def cascata_alt(*alts):
    """IF(a<>"",a_expr, IF(b<>"",b_expr, ... 0)) a partir de pares (nome, expressão)."""
    expr = '0'
    for nome, e in reversed(alts):
        expr = f'IF({cheio(nome)},{e},{expr})'
    return expr


# ── escrita do arquivo ──────────────────────────────────────────────────────
def escrever(form, caminho):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Formulário'
    ws.sheet_view.showGridLines = False
    for col, larg in zip('ABCD', (58, 30, 36, 46)):
        ws.column_dimensions[col].width = larg

    # cabeçalho fixo (4 linhas) ------------------------------------------------
    ws.merge_cells('A1:D1')
    c = ws['A1']
    c.value = form.titulo
    c.font = Font(bold=True, size=14, color='FFFFFF')
    c.fill = PatternFill('solid', fgColor=COR_TITULO)
    c.alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 26
    ws.merge_cells('A2:D2')
    c = ws['A2']
    c.value = form.subtitulo
    c.alignment = Alignment(wrap_text=True, vertical='top')
    c.font = Font(size=9)
    ws.row_dimensions[2].height = 92
    for col, txt in zip('ABCD', ('Campo', 'Valor', 'Unidade', 'Chave técnica (leitura por máquina)')):
        c = ws[f'{col}3']
        c.value = txt
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=COR_SECAO)
    ws.freeze_panes = 'A4'

    # 1ª passada: posição de cada linha ------------------------------------------
    inicio = 4
    pos = {}
    for i, ln in enumerate(form.linhas):
        pos[i] = inicio + i
    ref = {nome: f'$B${pos[i]}' for nome, i in form.nomes.items()}

    def resolve(txt):
        def sub(m):
            if m.group(1) not in ref:
                raise KeyError(f'referência desconhecida {{{m.group(1)}}}')
            return ref[m.group(1)]
        return re.sub(r'\{(\w+)\}', sub, txt)

    # listas auxiliares na coluna F (oculta), abaixo de tudo
    lista_pos = {}
    cursor = inicio + len(form.linhas) + 2
    for nome_l, itens in form.listas.items():
        lista_pos[nome_l] = cursor
        for k, it in enumerate(itens):
            ws.cell(row=cursor + k, column=6, value=it)
        cursor += len(itens) + 1
    if form.listas:
        ws.column_dimensions['F'].hidden = True

    dv_cache = {}

    def validacao(chave, **kw):
        if chave not in dv_cache:
            dv = DataValidation(allow_blank=True, **kw)
            ws.add_data_validation(dv)
            dv_cache[chave] = dv
        return dv_cache[chave]

    # 2ª passada: escrita ---------------------------------------------------------
    for i, ln in enumerate(form.linhas):
        row = pos[i]
        t = ln['t']
        if t in ('secao', 'grupo', 'nota'):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
            c = ws.cell(row=row, column=1, value=ln['texto'])
            if t == 'secao':
                cor = COR_CALC_SECAO if ln.get('calc') else COR_SECAO
                c.font = Font(bold=True, color='FFFFFF', size=11)
                c.fill = PatternFill('solid', fgColor=cor)
                ws.row_dimensions[row].height = 20
            elif t == 'grupo':
                c.font = Font(bold=True, italic=True)
                c.fill = PatternFill('solid', fgColor=COR_GRUPO)
            else:
                c.font = Font(italic=True, size=9, color='595959')
                c.alignment = Alignment(wrap_text=True, vertical='top')
                ws.row_dimensions[row].height = 13 * max(1, -(-len(ln['texto']) // 150))
            continue

        a = ws.cell(row=row, column=1, value=ln['label'])
        b = ws.cell(row=row, column=2)
        cu = ws.cell(row=row, column=3)
        d = ws.cell(row=row, column=4, value=ln['chave'])
        d.font = Font(size=8, color='7F7F7F')
        a.alignment = Alignment(wrap_text=True, vertical='center')
        cu.alignment = Alignment(wrap_text=True, vertical='center')
        d.alignment = Alignment(wrap_text=True, vertical='center')
        b.alignment = Alignment(horizontal='right', vertical='center', wrap_text=(t == 'campo'))
        b.number_format = ln['fmt']
        uni = ln['unidade']
        cu.value = resolve(uni) if isinstance(uni, str) and uni.startswith('=') else uni
        cu.font = Font(size=9)
        for c in (a, b, cu, d):
            c.border = BORDA

        if t == 'calc':
            b.value = '=' + resolve(ln['formula'])
            b.fill = PatternFill('solid', fgColor=COR_CALC)
            a.fill = PatternFill('solid', fgColor=COR_CALC)
            cu.fill = PatternFill('solid', fgColor=COR_CALC)
            d.fill = PatternFill('solid', fgColor=COR_CALC)
            b.font = Font(italic=True)
            continue

        # campo de entrada
        if ln['fixo'] is not None:
            b.value = ln['fixo']
            b.fill = PatternFill('solid', fgColor=COR_FIXO)
            b.font = Font(bold=True)
        else:
            b.fill = PatternFill('solid', fgColor=COR_ENTRADA)
            b.protection = Protection(locked=False)
            if ln['lista_ref']:
                itens = form.listas[ln['lista_ref']]
                ini = lista_pos[ln['lista_ref']]
                dv = validacao(('ref', ln['lista_ref'], ln['aviso']), type='list',
                               formula1=f'$F${ini}:$F${ini + len(itens) - 1}')
                if ln['aviso']:
                    dv.errorStyle = 'warning'
                dv.error = 'Valor fora da lista de referência — confirme.' if ln['aviso'] else 'Escolha um valor da lista.'
                dv.errorTitle = 'Confira o valor'
                dv.add(b)
            elif ln['lista']:
                dv = validacao(('lista', tuple(ln['lista'])), type='list',
                               formula1='"' + ','.join(ln['lista']) + '"')
                dv.error = 'Escolha um valor da lista.'
                dv.errorTitle = 'Valor inválido'
                dv.add(b)
            elif ln['maxlen']:
                dv = validacao(('len', ln['maxlen']), type='textLength', operator='lessThanOrEqual',
                               formula1=str(ln['maxlen']))
                dv.errorStyle = 'warning'
                dv.error = f'No máximo {ln["maxlen"]} caracteres.'
                dv.add(b)
            elif ln['minimo'] is not None or ln['maximo'] is not None:
                lo = ln['minimo'] if ln['minimo'] is not None else -1e9
                hi = ln['maximo'] if ln['maximo'] is not None else 1e12
                tipo = 'whole' if ln['inteiro'] else 'decimal'
                dv = validacao((tipo, lo, hi), type=tipo, operator='between',
                               formula1=str(lo), formula2=str(hi))
                dv.errorStyle = 'warning'
                dv.error = f'Valor fora da faixa esperada ({lo} a {hi}).'
                dv.errorTitle = 'Confira o valor'
                dv.add(b)

    # ── exemplo (só quando pedido por verificação) ──
    for nome, valor in form.exemplo.items():
        if nome in form.nomes:
            ws[ref[nome].replace('$', '')].value = valor

    ws.protection.sheet = True
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.sheet_properties.tabColor = COR_SECAO
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '3:3'
    wb.save(caminho)
    return {nome: ref[nome].replace('$', '') for nome in ref}


# ════════════════════════════════════════════════════════════════════════════
#  Blocos comuns
# ════════════════════════════════════════════════════════════════════════════
SN = ['Sim', 'Não']
LEGENDA = (
    'Como preencher: só as células AMARELAS da coluna B. Cinza-claro com valor em negrito = fixo do tipo de estudo '
    '(não editar); cinza com itálico = calculado pelo Excel, só para conferir (não editar). * = obrigatório. '
    'PERCENTUAL: digite com o símbolo % (4%), nunca só o número. O Excel guarda 0,04; a API do app espera pontos (4) e o importador converte pelo FORMATO da célula — por isso não troque o formato das células. '
    'Números sem separador de milhar e com ponto decimal quando digitados fora do Excel. '
    'CAMPO COM VÁRIAS UNIDADES: há uma linha por unidade — preencha UMA só e deixe as outras em branco; '
    'o Excel não converte uma na outra. Listas (produtos, custos…) têm vagas numeradas; vaga totalmente em branco = sem item.'
)


def bloco_criacao(f, tipo, nivel, ex):
    f.secao('1 · Criação do estudo')
    f.campo('nome', 'Nome do estudo *', 'texto (até 200)', 'estudos.nome', F_TXT, maxlen=200, ex=ex.get('nome'))
    f.campo('tipo', 'Tipo de empreendimento', 'fixo deste formulário', 'estudos.tipo_empreendimento', F_TXT, fixo=tipo)
    f.campo('nivel', 'Nível de análise', 'fixo deste formulário', 'estudos.nivel_analise', F_TXT, fixo=nivel)
    f.campo('origem', 'Origem do terreno *', 'lista: manual (inserir novo) | nucleo (buscar terreno — o importador só aceita manual)',
            'estudos.origem_terreno', F_TXT, lista=['manual', 'nucleo'], ex='manual')
    f.campo('uf', 'UF (recomendado)', 'texto, 2 letras', 'estudos.uf', F_TXT, maxlen=2, ex=ex.get('uf'))
    f.campo('notas', 'Notas', 'texto livre (só pela API)', 'estudos.notas', F_TXT, ex=ex.get('notas'))


def bloco_mercado(f, ex):
    f.secao('Análise de Mercado (opcional)')
    f.campo('regiao', 'Região de mercado', 'nome de região cadastrada na instância', 'estudos.regiao_mercado_id', F_TXT,
            ex=ex.get('regiao'))
    f.campo('apelo_tipo', 'Apelo comercial — tipo de dado', 'lista: anuncios | populacao | mercado | outro',
            'apelo_comercial_documentos.tipo_dado', F_TXT, lista=['anuncios', 'populacao', 'mercado', 'outro'],
            ex=ex.get('apelo_tipo'))
    f.campo('apelo_texto', 'Apelo comercial — texto adicional', 'texto livre (ou envie arquivo pdf/docx/xlsx à parte)',
            'apelo_comercial_documentos.texto_adicional', F_TXT, ex=ex.get('apelo_texto'))


# ════════════════════════════════════════════════════════════════════════════
#  PRELIMINAR
# ════════════════════════════════════════════════════════════════════════════
def preliminar(tipo, ex):
    """ex: dict nome->valor de exemplo (verificação); vazio na entrega."""
    lot = tipo == 'loteamento'
    rotulo = 'Loteamento' if lot else 'Incorporação'
    f = Form(f'Formulário de estudo — Preliminar · {rotulo}', LEGENDA)
    bloco_criacao(f, tipo, 'preliminar', ex)

    # ── Terreno & Áreas ──────────────────────────────────────────────────────
    f.secao('2 · Premissas › Terreno & Áreas')
    f.campo('terreno_nome', 'Nome do terreno (origem = manual)', 'texto (até 200)', 'estudos.terreno_manual_nome',
            F_TXT, maxlen=200, ex=ex.get('terreno_nome'))
    f.campo('terreno_area',
            'Área da poligonal / do terreno *' if lot else 'Área do terreno *',
            'm² (> 0)', 'estudos.terreno_manual_area', F_M2, minimo=0, ex=ex.get('terreno_area'))
    f.campo('nucleo_id', 'Imóvel do Núcleo (vínculo — só se origem = nucleo)', 'id do imóvel',
            'estudo_imoveis.imovel_nucleo_id', F_TXT)
    if lot:
        f.nota('Cascata de áreas do Loteamento. Cada linha aceita UMA entre as unidades abaixo (m², % da poligonal, '
               '% da parcelável). A APP não aceita % da parcelável (é ela que forma a parcelável).')
        linhas_area = [
            ('app', 'APP', False),
            ('elup', 'Espaço Livre de Uso Público (ELUP/EPU)', True),
            ('epc', 'EPC — Equipamento Público Comunitário', True),
            ('viapub', 'Sistema viário público', True),
            ('viapriv', 'Sistema viário privado', True),
            ('comuns', 'Áreas comuns privadas (incl. portaria)', True),
            ('verdes', 'Áreas verdes', True),
        ]
        chaves = dict(app='area_app', elup='area_elup_epu', epc='area_epc', viapub='area_viario_publico',
                      viapriv='area_viario_privado', comuns='area_comuns_privadas', verdes='area_verdes')
        for nome, rot, aceita_parc in linhas_area:
            f.grupo(f'{rot} — preencha uma unidade')
            k = chaves[nome]
            f.campo(f'{nome}_m2', f'{rot} — em m²', 'm²', f'{k}_modo = m2 → {k}_valor', F_M2, minimo=0,
                    ex=ex.get(f'{nome}_m2'))
            f.campo(f'{nome}_pp', f'{rot} — em % da poligonal', '% da área da poligonal',
                    f'{k}_modo = pct_poligonal → {k}_valor (em pontos: 8 = 8%)', F_PCT, minimo=0, maximo=1,
                    ex=ex.get(f'{nome}_pp'))
            if aceita_parc:
                f.campo(f'{nome}_pc', f'{rot} — em % da parcelável', '% da área parcelável',
                        f'{k}_modo = pct_parcelavel → {k}_valor (em pontos: 8 = 8%)', F_PCT, minimo=0, maximo=1,
                        ex=ex.get(f'{nome}_pc'))
    else:
        f.campo('coef_min', 'Coeficiente de aproveitamento mínimo', 'índice (x)', 'estudos.coef_aproveitamento_basico',
                F_X, minimo=0, ex=ex.get('coef_min'))
        f.campo('coef_max', 'Coeficiente de aproveitamento máximo', 'índice (x); teto = terreno × coeficiente',
                'estudos.coef_aproveitamento_maximo', F_X, minimo=0, ex=ex.get('coef_max'))
        f.campo('rf', 'Área privativa residencial FECHADA *', 'm² (base dos produtos residenciais)',
                'estudos.area_pvt_r_fechada', F_M2, minimo=0, ex=ex.get('rf'))
        f.campo('ra', 'Área privativa residencial aberta', 'm² (varandas etc.)', 'estudos.area_pvt_r_aberta', F_M2,
                minimo=0, ex=ex.get('ra'))
        f.campo('nf', 'Área privativa não residencial fechada (se houver NR)', 'm² (base dos produtos NR)',
                'estudos.area_pvt_nr_fechada', F_M2, minimo=0, ex=ex.get('nf'))
        f.campo('na', 'Área privativa não residencial aberta', 'm²', 'estudos.area_pvt_nr_aberta', F_M2, minimo=0,
                ex=ex.get('na'))
        f.campo('comum', 'Área comum total', 'm²', 'estudos.area_comum_total', F_M2, minimo=0, ex=ex.get('comum'))
    f.campo('imagem', 'Imagem principal (nome do arquivo; envio à parte)', 'png/jpeg/webp/pdf ≤ 25 MB',
            'estudo_documentos (categoria imagem_principal)', F_TXT)

    # ── Custos ───────────────────────────────────────────────────────────────
    f.secao('3 · Premissas › Custos')
    f.campo('cons_terreno', 'Considerar custo de aquisição do terreno *', 'Sim/Não (Não zera o custo)',
            'estudos.considerar_custo_terreno', F_TXT, lista=SN, ex='Sim')
    f.campo('cons_gestao', 'Considerar gestão e outros custos indiretos *', 'Sim/Não',
            'estudos.considerar_gestao_indiretos', F_TXT, lista=SN, ex='Sim')
    f.campo('cons_conting', 'Considerar contingências *', 'Sim/Não', 'estudos.considerar_contingencias', F_TXT,
            lista=SN, ex='Sim')
    f.campo('custo_terreno', 'Custo do terreno *', 'R$/m² de terreno', 'estudos.custo_terreno_m2', F_RS, minimo=0,
            ex=ex.get('custo_terreno'))

    if lot:
        f.grupo('Infraestrutura * — preencha uma unidade')
        f.campo('infra_pct', 'Infraestrutura — em % do VGV', '% do VGV', 'infra_modo = pct_vgv → infra_pct', F_PCT,
                minimo=0, maximo=10, ex=ex.get('infra_pct'))
        f.campo('infra_fixo', 'Infraestrutura — em R$ (valor fixo)', 'R$ (valor cheio)',
                'infra_modo = valor_fixo → infra_valor_fixo', F_RS, minimo=0, ex=ex.get('infra_fixo'))
        f.campo('infra_m2', 'Infraestrutura — em R$/m² de ALV', 'R$/m² de ALV (área líquida de venda)',
                'infra_modo = valor_m2 → custo_infra_m2', F_RS, minimo=0, ex=ex.get('infra_m2'))
    else:
        f.grupo('Construção * — preencha uma unidade')
        f.campo('constr_m2', 'Construção — em R$/m²', 'R$/m² de área privativa total',
                'construcao_modo = valor_m2 → custo_construcao_m2', F_RS, minimo=0, ex=ex.get('constr_m2'))
        f.campo('constr_total', 'Construção — em R$ (valor total)', 'R$ (valor cheio)',
                'construcao_modo = valor_total → construcao_valor_total', F_RS, minimo=0, ex=ex.get('constr_total'))

    f.grupo('Projetos * — preencha uma unidade')
    f.campo('proj_vgv', 'Projetos — em % do VGV', '% do VGV', 'projetos_modo = pct_vgv → projetos_pct', F_PCT,
            minimo=0, maximo=1, ex=ex.get('proj_vgv'))
    f.campo('proj_fixo', 'Projetos — em R$ (valor fixo)', 'R$ (valor cheio)',
            'projetos_modo = valor_fixo → projetos_valor_fixo', F_RS, minimo=0, ex=ex.get('proj_fixo'))
    f.campo('proj_constr', 'Projetos — em % do custo de obras',
            '% do custo de obras (infraestrutura)' if lot else '% do custo de obras (construção + decoração + gestão)',
            'projetos_modo = pct_constr → projetos_pct', F_PCT, minimo=0, maximo=1, ex=ex.get('proj_constr'))

    f.grupo('Demais custos (unidade fixa)')
    if not lot:
        f.campo('decoracao', 'Decoração', 'R$/m² de área privativa', 'estudos.custo_decoracao_m2', F_RS, minimo=0,
                ex=ex.get('decoracao'))
        f.campo('gestao_constr', 'Gestão da construção', '% de (construção + decoração)', 'estudos.taxa_gestao_pct',
                F_PCT, minimo=0, maximo=1, ex=ex.get('gestao_constr'))
        f.campo('incorp_reg', 'Incorporação e registro', '% do VGV', 'estudos.incorporacao_registro_pct', F_PCT,
                minimo=0, maximo=1, ex=ex.get('incorp_reg'))
    f.campo('manut', 'Manutenção pós-obra', '% do VGV', 'estudos.manutencao_pct', F_PCT, minimo=0, maximo=1,
            ex=ex.get('manut'))
    f.campo('conting', 'Contingências', '% do VGV', 'estudos.contingencias_pct', F_PCT, minimo=0, maximo=1,
            ex=ex.get('conting'))
    if lot:
        f.campo('stand', 'Stand de vendas', 'R$ (valor cheio)', 'estudos.stand_vendas_valor', F_RS, minimo=0,
                ex=ex.get('stand'))
    f.campo('gestao_ind', 'Gestão e outros custos indiretos', '% do VGV', 'estudos.gestao_indiretos_pct', F_PCT,
            minimo=0, maximo=1, ex=ex.get('gestao_ind'))
    f.campo('ret', 'Sujeito a RET *', 'Sim/Não (Sim = alíquota RET da instância)', 'estudos.sujeito_ret', F_TXT,
            lista=SN, ex=ex.get('ret'))
    f.campo('imposto', 'Imposto (se não RET)', '% do VGV', 'estudos.imposto_percentual', F_PCT, minimo=0, maximo=1,
            ex=ex.get('imposto'))
    f.campo('corretagem', 'Corretagem *', '% do VGV', 'estudos.corretagem_percentual', F_PCT, minimo=0, maximo=1,
            ex=ex.get('corretagem'))
    f.campo('marketing', 'Marketing *', '% do VGV', 'estudos.marketing_percentual', F_PCT, minimo=0, maximo=1,
            ex=ex.get('marketing'))

    # ── Permutas ─────────────────────────────────────────────────────────────
    f.secao('4 · Premissas › Permutas')
    cats = [('r', 'residencial', '')] if lot else [('r', 'residencial', ''), ('n', 'não residencial', '_nr')]
    for sigla, nome_cat, suf in cats:
        sufx = ' (residencial)' if lot else f' {nome_cat}'
        base_area = 'da área de venda (ALV)' if lot else ('da área de venda' if sigla == 'r' else 'da área de venda NR')
        f.grupo(f'Permuta física{sufx} — preencha uma unidade')
        f.campo(f'pf{sigla}_m2', f'Permuta física{sufx} — em m²', 'm²',
                f'permuta_fisica{suf}_modo = area_m2 → permuta_fisica{suf}_area_m2', F_M2, minimo=0,
                ex=ex.get(f'pf{sigla}_m2'))
        f.campo(f'pf{sigla}_pct', f'Permuta física{sufx} — em % da área de venda', '% ' + base_area,
                f'permuta_fisica{suf}_modo = pct_area_venda → permuta_fisica{suf}_pct', F_PCT, minimo=0, maximo=1,
                ex=ex.get(f'pf{sigla}_pct'))
        f.grupo(f'Permuta financeira{sufx} — preencha uma unidade')
        cat_k = 'residencial' if sigla == 'r' else 'nao_residencial'
        f.campo(f'pn{sigla}_pct', f'Permuta financeira{sufx} — em % do VGV', '% do VGV ' + ('' if lot else nome_cat),
                f'permuta_financeira_{cat_k}_modo = pct_vgv → permuta_financeira_{cat_k}_pct',
                F_PCT, minimo=0, maximo=1, ex=ex.get(f'pn{sigla}_pct'))
        f.campo(f'pn{sigla}_val', f'Permuta financeira{sufx} — em R$ (valor fixo)', 'R$ (valor cheio)',
                f'permuta_financeira_{cat_k}_modo = valor_fixo → permuta_financeira_{cat_k}_valor',
                F_RS, minimo=0, ex=ex.get(f'pn{sigla}_val'))

    bloco_mercado(f, ex)

    # ── Produtos ─────────────────────────────────────────────────────────────
    f.secao('5 · Produtos (Premissas › Produtos) — 10 vagas; Σ dos % = 100%')
    base_txt = 'ALV (área líquida de venda)' if lot else 'soma das áreas privativas fechadas (R + NR)'
    f.nota(f'O % de cada produto é sobre a {base_txt}. Área total, área média e VGV são calculados pelo app. '
           'Só entra no VGV a linha com % > 0, unidades > 0 e preço > 0.')
    N_PROD = 10
    for i in range(1, N_PROD + 1):
        p = f'prod{i}'
        f.grupo(f'Produto {i}')
        f.campo(f'{p}_nome', f'Produto {i} — Nome', 'texto (até 100)', f'preliminar_produtos[{i}].nome', F_TXT,
                maxlen=100, ex=ex.get(f'{p}_nome'))
        f.campo(f'{p}_tipo', f'Produto {i} — Tipo', 'lista: residencial | nao_residencial',
                f'preliminar_produtos[{i}].tipo', F_TXT, lista=['residencial', 'nao_residencial'], ex=ex.get(f'{p}_tipo'))
        f.campo(f'{p}_pct', f'Produto {i} — % da base', '% da base (Σ = 100%)', f'preliminar_produtos[{i}].pct_alv',
                F_PCT, minimo=0, maximo=1, ex=ex.get(f'{p}_pct'))
        f.campo(f'{p}_un', f'Produto {i} — Unidades', 'unidades (inteiro > 0)', f'preliminar_produtos[{i}].unidades',
                F_INT, minimo=0, inteiro=True, ex=ex.get(f'{p}_un'))
        f.campo(f'{p}_preco', f'Produto {i} — Preço de venda', 'R$/m² (> 0)' if lot else 'R$/m² privativo fechado (> 0)',
                f'preliminar_produtos[{i}].preco_venda_m2', F_RS, minimo=0, ex=ex.get(f'{p}_preco'))

    # ════════ referência calculada ════════
    f.secao('REFERÊNCIA CALCULADA — não preencher (espelha as fórmulas da Proforma do Preliminar)', calc=True)
    f.nota('Tudo abaixo é fórmula do Excel sobre o que foi preenchido acima, para conferir a ordem de grandeza. '
           'Quando a mesma grandeza foi preenchida em mais de uma unidade, vale a ordem: R$ fixo › R$/m² › % '
           '(e o alerta ao final acusa a duplicidade). Diferenças de centavos para o app são arredondamento.')
    f.calc('rp_ret', 'Alíquota RET da instância (suposição de referência)', '% da receita (parâmetro da instância)',
           '0.04', F_PCT, chave='premissa — conferir o parâmetro da instância')

    # áreas
    f.grupo('Áreas')
    if lot:
        f.calc('c_pol', 'Área da poligonal', 'm²', f'MAX(0,N({r("terreno_area")}))', F_M2)
        f.calc('c_app', 'APP', 'm²',
               f'MAX(0,{cascata_alt(("app_m2", r("app_m2")), ("app_pp", r("c_pol") + "*" + r("app_pp")))})', F_M2)
        f.calc('c_parc', 'Área parcelável', 'm²', f'MAX(0,{r("c_pol")}-{r("c_app")})', F_M2)

        def linha_cascata(nome, rot):
            f.calc(f'c_{nome}', rot, 'm²',
                   'MAX(0,' + cascata_alt((f'{nome}_m2', r(f'{nome}_m2')),
                                          (f'{nome}_pp', r('c_pol') + '*' + r(f'{nome}_pp')),
                                          (f'{nome}_pc', r('c_parc') + '*' + r(f'{nome}_pc'))) + ')', F_M2)
        for nome, rot in (('elup', 'ELUP/EPU'), ('epc', 'EPC'), ('viapub', 'Sistema viário público')):
            linha_cascata(nome, rot)
        f.calc('c_liq', 'Área líquida', 'm²', f'MAX(0,{r("c_parc")}-({soma("c_elup","c_epc","c_viapub")}))', F_M2)
        for nome, rot in (('viapriv', 'Sistema viário privado'), ('comuns', 'Áreas comuns privadas'),
                          ('verdes', 'Áreas verdes')):
            linha_cascata(nome, rot)
        f.calc('c_alv', 'Área líquida de venda (ALV) = base dos produtos', 'm²',
               f'MAX(0,{r("c_liq")}-({soma("c_viapriv","c_comuns","c_verdes")}))', F_M2)
        f.calc('c_alv_ha', 'ALV', 'hectares', f'{r("c_alv")}/10000', F_M2)
        f.calc('c_alv_pct', 'ALV / poligonal ("vendável / gleba")', '% da poligonal',
               f'IF({r("c_pol")}>0,{r("c_alv")}/{r("c_pol")},"")', F_PCT)
        base = 'c_alv'
        T = 'c_pol'
        vend = 'c_alv'
    else:
        f.calc('c_T', 'Área do terreno', 'm²', f'MAX(0,N({r("terreno_area")}))', F_M2)
        f.calc('c_priv', 'Área privativa total (4 áreas privativas)', 'm²',
               f'MAX(0,N({r("rf")}))+MAX(0,N({r("ra")}))+MAX(0,N({r("nf")}))+MAX(0,N({r("na")}))', F_M2)
        f.calc('c_constr', 'Área construída total (privativa + comum)', 'm²', f'{r("c_priv")}+MAX(0,N({r("comum")}))', F_M2)
        f.calc('c_vend', 'Área vendável = base dos produtos (privativas fechadas R + NR)', 'm²',
               f'MAX(0,N({r("rf")}))+MAX(0,N({r("nf")}))', F_M2)
        f.calc('c_teto', 'Teto de aproveitamento (terreno × coeficiente máximo)', 'm²',
               f'IF(N({r("coef_max")})>0,{r("c_T")}*{r("coef_max")},"")', F_M2)
        f.calc('c_aprov', 'Uso do teto (área privativa total ÷ teto)', '% do teto',
               f'IF(ISNUMBER({r("c_teto")}),IF({r("c_teto")}>0,{r("c_priv")}/{r("c_teto")},""),"")', F_PCT)
        f.calc('c_efic', 'Eficiência (área vendável ÷ terreno)', '% do terreno',
               f'IF({r("c_T")}>0,{r("c_vend")}/{r("c_T")},"")', F_PCT)
        base = 'c_vend'
        T = 'c_T'
        vend = 'c_vend'

    # catálogo
    f.grupo('Produtos (catálogo efetivo = % > 0, unidades > 0, preço > 0)')
    for i in range(1, N_PROD + 1):
        p = f'prod{i}'
        eff = f'AND({r(base)}*N({r(p+"_pct")})>0,N({r(p+"_un")})>0,N({r(p+"_preco")})>0)'
        f.calc(f'{p}_area', f'Produto {i} — área total (base × %)', 'm²',
               f'IF({eff},{r(base)}*{r(p+"_pct")},0)', F_M2)
        f.calc(f'{p}_amed', f'Produto {i} — área média por unidade', 'm²',
               f'IF({r(p+"_area")}>0,{r(p+"_area")}/{r(p+"_un")},"")', F_M2)
        f.calc(f'{p}_vgv', f'Produto {i} — VGV bruto (área total × preço)', 'R$',
               f'{r(p+"_area")}*N({r(p+"_preco")})', F_RS)
    f.calc('c_spct', 'Σ % da base dos produtos (deve ser 100%)', '%',
           'SUM(' + ','.join(r(f'prod{i}_pct') for i in range(1, N_PROD + 1)) + ')', F_PCT)

    def por_tipo(campo_valor, tipo_alvo):
        itens = range(1, N_PROD + 1)
        if lot:
            return '+'.join(r(f'prod{i}_{campo_valor}') for i in itens)
        if tipo_alvo == 'nao_residencial':
            return '+'.join(f'IF({r(f"prod{i}_tipo")}="nao_residencial",{r(f"prod{i}_{campo_valor}")},0)' for i in itens)
        return '+'.join(f'IF({r(f"prod{i}_tipo")}<>"nao_residencial",{r(f"prod{i}_{campo_valor}")},0)' for i in itens)

    f.calc('c_vgvb_r', 'VGV bruto' + ('' if lot else ' residencial'), 'R$', por_tipo('vgv', 'residencial'), F_RS)
    f.calc('c_area_r', 'Área do catálogo' + ('' if lot else ' residencial'), 'm²', por_tipo('area', 'residencial'), F_M2)
    if not lot:
        f.calc('c_vgvb_n', 'VGV bruto não residencial', 'R$', por_tipo('vgv', 'nao_residencial'), F_RS)
        f.calc('c_area_n', 'Área do catálogo não residencial', 'm²', por_tipo('area', 'nao_residencial'), F_M2)
    f.calc('c_unid', 'Nº de unidades do catálogo (linhas efetivas)', 'unidades',
           '+'.join(f'IF({r(f"prod{i}_area")}>0,N({r(f"prod{i}_un")}),0)' for i in range(1, N_PROD + 1)), F_INT)

    # permutas
    f.grupo('Permutas')
    base_perm_r = 'c_alv' if lot else 'c_area_r'
    f.calc('c_pf_r_area', 'Permuta física' + (' ' if lot else ' residencial ') + '— área entregue', 'm²',
           'MAX(0,' + cascata_alt(('pfr_m2', r('pfr_m2')), ('pfr_pct', r(base_perm_r) + '*' + r('pfr_pct'))) + ')', F_M2)
    f.calc('c_pm_r', 'Preço médio do catálogo' + ('' if lot else ' residencial'), 'R$/m²',
           f'IF({r("c_area_r")}>0,{r("c_vgvb_r")}/{r("c_area_r")},0)', F_RS)
    f.calc('c_pf_r_vgv', 'VGV da permuta física' + ('' if lot else ' residencial') + ' (limitado ao VGV bruto)', 'R$',
           f'MIN({r("c_pf_r_area")}*{r("c_pm_r")},{r("c_vgvb_r")})', F_RS)
    if not lot:
        f.calc('c_pf_n_area', 'Permuta física não residencial — área entregue', 'm²',
               'MAX(0,' + cascata_alt(('pfn_m2', r('pfn_m2')), ('pfn_pct', r('c_area_n') + '*' + r('pfn_pct'))) + ')',
               F_M2)
        f.calc('c_pm_n', 'Preço médio do catálogo não residencial', 'R$/m²',
               f'IF({r("c_area_n")}>0,{r("c_vgvb_n")}/{r("c_area_n")},0)', F_RS)
        f.calc('c_pf_n_vgv', 'VGV da permuta física não residencial (limitado ao VGV bruto)', 'R$',
               f'MIN({r("c_pf_n_area")}*{r("c_pm_n")},{r("c_vgvb_n")})', F_RS)
    f.calc('c_vgv_r', 'VGV' + ('' if lot else ' residencial') + ' (bruto − permuta física)', 'R$',
           f'{r("c_vgvb_r")}-{r("c_pf_r_vgv")}', F_RS)
    if not lot:
        f.calc('c_vgv_n', 'VGV não residencial (bruto − permuta física)', 'R$', f'{r("c_vgvb_n")}-{r("c_pf_n_vgv")}', F_RS)
    f.calc('c_vgv', 'VGV (receita bruta do incorporador)' if not lot else 'VGV (receita bruta)', 'R$',
           r('c_vgv_r') + ('' if lot else '+' + r('c_vgv_n')), F_RS)
    areafis = r('c_pf_r_area') + ('' if lot else '+' + r('c_pf_n_area'))
    f.calc('c_areafis', 'Área total permutada fisicamente', 'm²', areafis, F_M2)
    f.calc('c_vend_liq', 'Área vendável líquida (vendável − permuta física)', 'm²', f'{r(vend)}-{r("c_areafis")}', F_M2)
    f.calc('c_pfin_r', 'Permuta financeira' + ('' if lot else ' residencial'), 'R$',
           f'IF({cheio("pnr_val")},{r("pnr_val")},{r("c_vgv_r")}*N({r("pnr_pct")}))', F_RS)
    pfin = r('c_pfin_r')
    if not lot:
        f.calc('c_pfin_n', 'Permuta financeira não residencial', 'R$',
               f'IF({cheio("pnn_val")},{r("pnn_val")},{r("c_vgv_n")}*N({r("pnn_pct")}))', F_RS)
        pfin += '+' + r('c_pfin_n')

    # deduções
    f.grupo('Deduções da receita')
    f.calc('c_imp_pct', 'Alíquota de imposto aplicada', '% do VGV',
           f'IF({r("ret")}="Sim",{r("rp_ret")},N({r("imposto")}))', F_PCT)
    f.calc('c_imp', 'Imposto', 'R$', f'{r("c_vgv")}*{r("c_imp_pct")}', F_RS)
    f.calc('c_corr', 'Corretagem', 'R$', f'{r("c_vgv")}*N({r("corretagem")})', F_RS)
    f.calc('c_mkt', 'Marketing', 'R$', f'{r("c_vgv")}*N({r("marketing")})', F_RS)
    f.calc('c_rliq', 'Receita líquida (VGV − deduções)', 'R$',
           f'{r("c_vgv")}-{r("c_imp")}-{r("c_corr")}-{r("c_mkt")}-({pfin})', F_RS)

    # custos diretos
    f.grupo('Custos diretos')
    f.calc('c_cterr', 'Custo do terreno', 'R$',
           f'IF({r("cons_terreno")}="Não",0,N({r("custo_terreno")})*{r(T)})', F_RS)
    if lot:
        f.calc('c_infra', 'Infraestrutura', 'R$',
               'IF(' + cheio('infra_fixo') + ',' + r('infra_fixo') + ',IF(' + cheio('infra_m2') + ',' +
               r('infra_m2') + '*' + r('c_alv') + ',' + r('c_vgv') + '*N(' + r('infra_pct') + ')))', F_RS)
        f.calc('c_obras', 'Custo de obras (= infraestrutura)', 'R$', r('c_infra'), F_RS)
    else:
        f.calc('c_constr_rs', 'Construção', 'R$',
               f'IF({cheio("constr_total")},{r("constr_total")},N({r("constr_m2")})*{r("c_priv")})', F_RS)
        f.calc('c_decor', 'Decoração', 'R$', f'N({r("decoracao")})*{r("c_priv")}', F_RS)
        f.calc('c_gestc', 'Gestão da construção', 'R$', f'({r("c_constr_rs")}+{r("c_decor")})*N({r("gestao_constr")})', F_RS)
        f.calc('c_obras', 'Custo de obras (construção + decoração + gestão)', 'R$',
               soma('c_constr_rs', 'c_decor', 'c_gestc'), F_RS)
    f.calc('c_proj', 'Projetos', 'R$',
           f'IF({cheio("proj_fixo")},{r("proj_fixo")},IF({cheio("proj_constr")},{r("c_obras")}*{r("proj_constr")},'
           f'{r("c_vgv")}*N({r("proj_vgv")})))', F_RS)
    if not lot:
        f.calc('c_increg', 'Incorporação e registro', 'R$', f'{r("c_vgv")}*N({r("incorp_reg")})', F_RS)
    f.calc('c_manut', 'Manutenção pós-obra', 'R$', f'{r("c_vgv")}*N({r("manut")})', F_RS)
    f.calc('c_cont', 'Contingências', 'R$', f'IF({r("cons_conting")}="Não",0,{r("c_vgv")}*N({r("conting")}))', F_RS)
    diretos = ['c_cterr', 'c_proj', 'c_manut', 'c_cont'] + (['c_infra'] if lot else ['c_constr_rs', 'c_gestc', 'c_decor', 'c_increg'])
    f.calc('c_cdir', 'Custo direto total', 'R$', soma(*diretos), F_RS)

    # indiretos e resultado
    f.grupo('Custos indiretos e resultado')
    if lot:
        f.calc('c_stand', 'Stand de vendas', 'R$', f'N({r("stand")})', F_RS)
    f.calc('c_gind', 'Gestão e outros custos indiretos', 'R$',
           f'IF({r("cons_gestao")}="Não",0,{r("c_vgv")}*N({r("gestao_ind")}))', F_RS)
    f.calc('c_cind', 'Custo indireto total', 'R$', soma(*((['c_stand'] if lot else []) + ['c_gind'])), F_RS)
    f.calc('c_rop', 'Receita operacional (receita líquida − custo direto)', 'R$', f'{r("c_rliq")}-{r("c_cdir")}', F_RS)
    f.calc('c_res', 'RESULTADO', 'R$', f'{r("c_rop")}-{r("c_cind")}', F_RS)
    f.calc('c_marg', 'Margem sobre VGV', '% do VGV', f'IF({r("c_vgv")}>0,{r("c_res")}/{r("c_vgv")},"")', F_PCT)
    f.calc('c_obras_pct', 'Custo de obras ÷ VGV', '% do VGV', f'IF({r("c_vgv")}>0,{r("c_obras")}/{r("c_vgv")},"")', F_PCT)
    f.calc('c_rliq_pct', 'Receita líquida ÷ VGV', '% do VGV', f'IF({r("c_vgv")}>0,{r("c_rliq")}/{r("c_vgv")},"")', F_PCT)
    f.calc('c_inv', 'Investimento total (direto + indireto)', 'R$', f'{r("c_cdir")}+{r("c_cind")}', F_RS)
    f.calc('c_roi', 'ROI (resultado ÷ investimento)', '%', f'IF({r("c_inv")}>0,{r("c_res")}/{r("c_inv")},"")', F_PCT)
    f.calc('c_pmu', 'Preço médio por unidade', 'R$', f'IF({r("c_unid")}>0,{r("c_vgv")}/{r("c_unid")},"")', F_RS)

    # equivalências de unidade (referência)
    f.grupo('Equivalências de unidade (referência) — mesmo valor em outras unidades')
    if lot:
        f.calc('e_infra_m2', 'Infraestrutura em R$/m² de ALV', 'R$/m² de ALV',
               f'IF({r("c_alv")}>0,{r("c_infra")}/{r("c_alv")},"")', F_RS)
        f.calc('e_infra_pct', 'Infraestrutura em % do VGV', '% do VGV',
               f'IF({r("c_vgv")}>0,{r("c_infra")}/{r("c_vgv")},"")', F_PCT)
    else:
        f.calc('e_constr_m2', 'Construção em R$/m² de área privativa', 'R$/m² priv. total',
               f'IF({r("c_priv")}>0,{r("c_constr_rs")}/{r("c_priv")},"")', F_RS)
        f.calc('e_constr_vgv', 'Construção em % do VGV', '% do VGV',
               f'IF({r("c_vgv")}>0,{r("c_constr_rs")}/{r("c_vgv")},"")', F_PCT)
    f.calc('e_proj_pct', 'Projetos em % do VGV', '% do VGV', f'IF({r("c_vgv")}>0,{r("c_proj")}/{r("c_vgv")},"")', F_PCT)
    f.calc('e_terr_total', 'Custo do terreno em R$ por m² vendável', 'R$/m² vendável',
           f'IF({r(vend)}>0,{r("c_cterr")}/{r(vend)},"")', F_RS)
    f.calc('e_pf_pct', 'Permuta física em % da área de venda', '% da área de venda',
           f'IF({r(vend)}>0,{r("c_areafis")}/{r(vend)},"")', F_PCT)
    f.calc('e_pfin_pct', 'Permuta financeira total em % do VGV', '% do VGV',
           f'IF({r("c_vgv")}>0,({pfin})/{r("c_vgv")},"")', F_PCT)

    # alertas
    f.grupo('Conferências')
    f.calc('a_pct', 'Σ % dos produtos', 'alerta',
           f'IF({r("c_spct")}=0,"sem produtos",IF(ABS({r("c_spct")}-1)<0.00005,"OK","⚠ Σ ≠ 100% (o app recusa salvar)"))')
    dup = []

    def dupl(nome, rot, *alts):
        f.calc(nome, f'Unidades em duplicidade — {rot}', 'alerta',
               f'IF(COUNTA({",".join(r(a) for a in alts)})>1,"⚠ mais de uma unidade preenchida","OK")')
    if lot:
        for nm, rot in (('app', 'APP'), ('elup', 'ELUP/EPU'), ('epc', 'EPC'), ('viapub', 'Viário público'),
                        ('viapriv', 'Viário privado'), ('comuns', 'Comuns privadas'), ('verdes', 'Verdes')):
            alts = [f'{nm}_m2', f'{nm}_pp'] + ([] if nm == 'app' else [f'{nm}_pc'])
            dupl(f'a_d_{nm}', rot, *alts)
        dupl('a_d_infra', 'Infraestrutura', 'infra_pct', 'infra_fixo', 'infra_m2')
    else:
        dupl('a_d_constr', 'Construção', 'constr_m2', 'constr_total')
    dupl('a_d_proj', 'Projetos', 'proj_vgv', 'proj_fixo', 'proj_constr')
    dupl('a_d_pfr', 'Permuta física' + ('' if lot else ' R'), 'pfr_m2', 'pfr_pct')
    dupl('a_d_pnr', 'Permuta financeira' + ('' if lot else ' R'), 'pnr_pct', 'pnr_val')
    if not lot:
        dupl('a_d_pfn', 'Permuta física NR', 'pfn_m2', 'pfn_pct')
        dupl('a_d_pnn', 'Permuta financeira NR', 'pnn_pct', 'pnn_val')
        f.calc('a_aprov', 'Aproveitamento do coeficiente', 'alerta',
               f'IF(ISNUMBER({r("c_aprov")}),IF({r("c_aprov")}>1,"⚠ área privativa acima do teto","OK"),"sem coeficiente máximo")')
    cond = f'{r("c_pf_r_area")}*{r("c_pm_r")}>{r("c_vgvb_r")}+0.005'
    if not lot:
        cond = f'OR({cond},{r("c_pf_n_area")}*{r("c_pm_n")}>{r("c_vgvb_n")}+0.005)'
    f.calc('a_permuta', 'Permuta física limitada ao VGV bruto', 'alerta',
           f'IF({cond},"⚠ permuta acima da base: o app corta no VGV bruto","OK")')
    return f



# ════════════════════════════════════════════════════════════════════════════
#  AVANÇADO
# ════════════════════════════════════════════════════════════════════════════
CATEGORIAS_CUSTO = ['Preço', 'Registro', 'Outro', 'Construção', 'Outorga', 'Decoração', 'Gestão da obra',
                    'Contingência', 'Marketing & Publicidade', 'Corretagem de vendas', 'Projetos',
                    'Licenças e Aprovações', 'Manutenção pós-obra', 'Marketing global', 'Gestão',
                    'Stand de vendas', 'Juros de financiamento', 'Taxas bancárias', 'Estruturação de dívida',
                    'Investidores']
CRONOGRAMA_CUSTO = ['planejamento', 'pre_lancamento', 'lancamento', 'obra', 'pos_obra', 'customizado']
UNIDADES_CUSTO = ['rs', 'rs_m2_priv', 'rs_m2_terreno', 'pct_vgv', 'pct_obra', 'pct_recebido', 'pct_constr']
ROTULO_UNIDADE = (
    '=IF({u}="","escolha a Unidade acima",IF({u}="rs","R$ (valor cheio)",IF({u}="rs_m2_priv","R$/m² privativo (fechada+aberta)",'
    'IF({u}="rs_m2_terreno","R$/m² de terreno",IF(LEFT({u},3)="pct","— (use a linha de %)","unidade inválida")))))'
)
ROTULO_UNIDADE_PCT = (
    '=IF({u}="","escolha a Unidade acima",IF({u}="pct_vgv","% do VGV",IF({u}="pct_receita","% da receita líquida (RET, sem juros)",'
    'IF({u}="pct_obra","% da obra (grupo Obras inteiro)",IF({u}="pct_recebido","% do recebido (com juros de tabela)",'
    'IF({u}="pct_constr","% da construção (constr.+decor.+gestão da obra)",IF(LEFT({u},2)="rs","— (use a linha de R$)","unidade inválida"))))))'
)

# Unidades aceitas por (grupo, categoria) — cópia da tabela da tela de Custos; a tela é a autoridade.
UNIDADES_POR_CATEGORIA = {
    ('terreno', 'Preço'): ['rs', 'rs_m2_terreno'], ('terreno', 'Registro'): ['rs', 'rs_m2_priv'], ('terreno', 'Outro'): ['rs', 'pct_vgv'],
    ('obra', 'Construção'): ['rs', 'rs_m2_priv'], ('obra', 'Outorga'): ['rs', 'pct_vgv'], ('obra', 'Decoração'): ['rs', 'rs_m2_priv'],
    ('obra', 'Gestão da obra'): ['rs', 'pct_obra'], ('obra', 'Contingência'): ['rs', 'pct_vgv'], ('obra', 'Outro'): ['rs', 'pct_vgv'],
    ('diretos', 'Marketing & Publicidade'): ['rs', 'pct_vgv', 'pct_recebido'], ('diretos', 'Corretagem de vendas'): ['pct_vgv'],
    ('diretos', 'Projetos'): ['rs', 'rs_m2_priv', 'pct_constr'], ('diretos', 'Licenças e Aprovações'): ['rs', 'rs_m2_priv'],
    ('diretos', 'Manutenção pós-obra'): ['rs', 'pct_vgv', 'pct_recebido'], ('diretos', 'Outro'): ['rs', 'pct_vgv', 'pct_recebido'],
    ('indireto', 'Marketing global'): ['rs', 'pct_vgv', 'pct_recebido'], ('indireto', 'Stand de vendas'): ['rs', 'pct_vgv'],
    ('indireto', 'Gestão'): ['rs', 'pct_vgv', 'pct_recebido'], ('indireto', 'Outro'): ['rs', 'pct_vgv', 'pct_recebido'],
    ('financeiro', 'Juros de financiamento'): ['rs', 'pct_vgv'], ('financeiro', 'Taxas bancárias'): ['rs', 'pct_vgv'],
    ('financeiro', 'Estruturação de dívida'): ['rs', 'pct_vgv'], ('financeiro', 'Investidores'): ['rs', 'pct_vgv'],
    ('financeiro', 'Outro'): ['rs', 'pct_vgv'],
}

NOTA_CUSTOS = (
    'Categorias e unidades aceitas (rs = R$ · rs_m2_priv = R$/m² privativo · rs_m2_terreno · pct_vgv · pct_obra · '
    'pct_recebido · pct_constr). TERRENO: Preço (rs, rs_m2_terreno; permuta financeira: rs, pct_vgv) · Registro (rs, rs_m2_priv) · Outro (rs, pct_vgv). '
    'OBRA: Construção, Decoração (rs, rs_m2_priv) · Outorga, Contingência, Outro (rs, pct_vgv) · Gestão da obra (rs, pct_obra). '
    'DIRETOS: Marketing & Publicidade, Manutenção pós-obra, Outro (rs, pct_vgv, pct_recebido) · Corretagem de vendas (pct_vgv) · '
    'Projetos (rs, rs_m2_priv, pct_constr) · Licenças e Aprovações (rs, rs_m2_priv). INDIRETO: Marketing global, Gestão, Outro '
    '(rs, pct_vgv, pct_recebido) · Stand de vendas (rs, pct_vgv). FINANCEIRO: Juros de financiamento, Taxas bancárias, Estruturação de dívida, '
    'Investidores, Outro (rs, pct_vgv). Cronograma padrão: terreno/Preço = customizado; terreno/demais, indireto = planejamento; '
    'obra/diretos/financeiro = obra; Manutenção pós-obra = pos_obra. Início e Duração só valem com cronograma = customizado; '
    'cronograma ancorado em fase customizada = escrever "fase: <nome da fase>". Corretagem sai no mês da venda e a permuta física '
    'na entrega das unidades (não escolhem distribuição). Obrigatórias: terreno/Preço, obra/Construção, diretos/Corretagem de vendas (% VGV).'
)


def avancado(tipo, ex):
    lot = tipo == 'loteamento'
    rotulo = 'Loteamento' if lot else 'Incorporação'
    f = Form(f'Formulário de estudo — Avançado · {rotulo}', LEGENDA + ' Meses: 0 = mês da data de início do projeto.')
    f.listas['categorias'] = CATEGORIAS_CUSTO
    f.listas['cronograma'] = CRONOGRAMA_CUSTO
    f.listas['distribuicao'] = ['Linear', 'Curva S', 'unit_delivery', 'sales_revenue']
    bloco_criacao(f, tipo, 'avancado', ex)

    # ── Empreendimento › Informações ──
    f.secao('2 · Empreendimento › Informações')
    f.campo('matricula', 'Matrícula', 'texto (até 100)', 'estudos.matricula', F_TXT, maxlen=100, ex=ex.get('matricula'))
    f.campo('descricao', 'Descrição', 'texto longo', 'estudos.descricao', F_TXT, ex=ex.get('descricao'))
    f.campo('terreno_nome', 'Nome do terreno (origem = manual)', 'texto (até 200)', 'estudos.terreno_manual_nome',
            F_TXT, maxlen=200, ex=ex.get('terreno_nome'))
    f.campo('terreno_area', 'Área do terreno *', 'm² (> 0; base de "R$/m² de terreno")', 'estudos.terreno_manual_area',
            F_M2, minimo=0, ex=ex.get('terreno_area'))
    f.campo('nucleo_id', 'Imóvel do Núcleo (vínculo — só se origem = nucleo)', 'id do imóvel',
            'estudo_imoveis.imovel_nucleo_id', F_TXT)
    if not lot:
        f.campo('coef_min', 'Coeficiente de aproveitamento mínimo (informativo)', 'índice (x)',
                'estudos.coef_aproveitamento_basico', F_X, minimo=0, ex=ex.get('coef_min'))
        f.campo('coef_max', 'Coeficiente de aproveitamento máximo (informativo)', 'índice (x)',
                'estudos.coef_aproveitamento_maximo', F_X, minimo=0, ex=ex.get('coef_max'))
    f.campo('imagem', 'Imagem principal / render / planta (nome do arquivo; envio à parte)', 'png/jpeg/webp/pdf ≤ 25 MB',
            'estudo_documentos', F_TXT)

    # ── Financeiro ──
    f.secao('3 · Viabilidade › Financeiro e tributação')
    f.campo('juros_tabela', 'Juros de tabela *', '% a.a. (0 a 999,99; vale p/ todos os grupos)',
            'estudos.juros_tabela_aa_padrao', F_PCT, minimo=0, maximo=9.9999, ex=ex.get('juros_tabela'))
    f.campo('taxa_desc', 'Taxa de desconto p/ VP *', '% a.a. (0 a 100)', 'estudos.taxa_desconto_aa', F_PCT,
            minimo=0, maximo=1, ex=ex.get('taxa_desc'))
    f.campo('corr_perm', 'Corretagem incide sobre a permuta física *', 'Sim/Não',
            'estudos.corretagem_sobre_permuta_fisica', F_TXT, lista=SN, ex=ex.get('corr_perm'))
    f.campo('ret', 'RET (Regime Especial de Tributação) *', 'Sim/Não (único imposto aplicado pelo Avançado)',
            'parametros.considerar_ret', F_TXT, lista=SN, ex=ex.get('ret'))
    f.campo('ret_pct', 'Alíquota RET *', '% da receita (0 a 100; em loteamento use a alíquota efetiva)',
            'parametros.ret_pct', F_PCT, minimo=0, maximo=1, ex=ex.get('ret_pct'))

    # ── Cronograma ──
    f.secao('4 · Cronograma (meses a partir do mês 0)')
    f.nota('Pré-lançamento, Lançamento e Pós-obras começam no FIM do evento anterior (travado no app): esses inícios estão na '
           'zona calculada, não são preenchidos. Informe "tem Pré-lançamento" ANTES de preencher o cronograma.')
    f.campo('data_ini', 'Data de início do projeto *', 'data (mês/ano) — mês 0 do fluxo', 'parametros.data_inicio_projeto',
            F_DATA, ex=ex.get('data_ini'))
    f.campo('tem_pre', 'Este empreendimento tem Pré-lançamento *', 'Sim/Não', 'parametros.tem_pre_lancamento', F_TXT,
            lista=SN, ex=ex.get('tem_pre'))
    f.campo('plan_ini', 'Planejamento — Início *', 'mês (inteiro ≥ 0; 0 = início)', 'cronograma.planejamento.inicio_mes',
            F_INT, minimo=0, inteiro=True, ex=ex.get('plan_ini'))
    f.campo('plan_dur', 'Planejamento — Duração *', 'meses (inteiro ≥ 1)', 'cronograma.planejamento.duracao_meses',
            F_INT, minimo=1, inteiro=True, ex=ex.get('plan_dur'))
    f.campo('pre_dur', 'Pré-lançamento — Duração (só se tiver pré)', 'meses (inteiro ≥ 1)',
            'cronograma.pre_lancamento.duracao_meses', F_INT, minimo=1, inteiro=True, ex=ex.get('pre_dur'))
    f.campo('lanc_dur', 'Lançamento — Duração *', 'meses (inteiro ≥ 1)', 'cronograma.lancamento.duracao_meses',
            F_INT, minimo=1, inteiro=True, ex=ex.get('lanc_dur'))
    f.campo('obra_ini', 'Obra — Início *', 'mês (inteiro ≥ 0)', 'cronograma.obra.inicio_mes', F_INT, minimo=0,
            inteiro=True, ex=ex.get('obra_ini'))
    f.campo('obra_dur', 'Obra — Duração *', 'meses (inteiro ≥ 1)', 'cronograma.obra.duracao_meses', F_INT,
            minimo=1, inteiro=True, ex=ex.get('obra_dur'))
    f.campo('pos_dur', 'Pós-obras — Duração *', 'meses (inteiro ≥ 1)', 'cronograma.pos_obra.duracao_meses', F_INT,
            minimo=1, inteiro=True, ex=ex.get('pos_dur'))

    f.secao('Fases customizadas do cronograma (opcional) — 5 vagas; servem de âncora para custos e funding')
    N_FASE = 5
    for i in range(1, N_FASE + 1):
        f.campo(f'fase{i}_nome', f'Fase {i} — Nome', 'texto', f'fases[{i}].nome', F_TXT,
                ex=ex.get(f'fase{i}_nome'))
        f.campo(f'fase{i}_ini', f'Fase {i} — Início', 'mês (inteiro ≥ 0)', f'fases[{i}].inicio_mes', F_INT,
                minimo=0, inteiro=True, ex=ex.get(f'fase{i}_ini'))
        f.campo(f'fase{i}_dur', f'Fase {i} — Duração', 'meses (inteiro ≥ 1)', f'fases[{i}].duracao_meses',
                F_INT, minimo=1, inteiro=True, ex=ex.get(f'fase{i}_dur'))

    bloco_mercado(f, ex)

    # ── Tipologias / Lotes ──
    N_TIP = 10
    nome_item = 'Lote' if lot else 'Tipologia'
    f.secao(f'5 · {"Lotes" if lot else "Tipologias"} (Empreendimento › Tipologias) — {N_TIP} vagas')
    if lot:
        f.nota('Em Loteamento o tipo é sempre "lote"; a tela não mostra Tipo, Dormitórios e Vagas.')
    for i in range(1, N_TIP + 1):
        t = f'tip{i}'
        f.grupo(f'{nome_item} {i}')
        f.campo(f'{t}_nome', f'{nome_item} {i} — Nome', 'texto (até 100)', f'tipologias[{i}].nome', F_TXT,
                maxlen=100, ex=ex.get(f'{t}_nome'))
        if not lot:
            f.campo(f'{t}_tipo', f'Tipologia {i} — Tipo', 'lista: apartamento | cobertura | loja | outro',
                    f'tipologias[{i}].tipo_unidade', F_TXT, lista=['apartamento', 'cobertura', 'loja', 'outro'],
                    ex=ex.get(f'{t}_tipo'))
        f.campo(f'{t}_priv', f'{nome_item} {i} — Área privativa (fechada)', 'm²', f'tipologias[{i}].area_privativa_m2',
                F_M2, minimo=0, ex=ex.get(f'{t}_priv'))
        f.campo(f'{t}_aberta', f'{nome_item} {i} — Área aberta', 'm² (entra no VGV)',
                f'tipologias[{i}].area_privativa_aberta_m2', F_M2, minimo=0, ex=ex.get(f'{t}_aberta'))
        if not lot:
            f.campo(f'{t}_dorm', f'Tipologia {i} — Dormitórios', 'inteiro', f'tipologias[{i}].dormitorios', F_INT,
                    minimo=0, inteiro=True, ex=ex.get(f'{t}_dorm'))
            f.campo(f'{t}_vagas', f'Tipologia {i} — Vagas', 'inteiro', f'tipologias[{i}].vagas', F_INT,
                    minimo=0, inteiro=True, ex=ex.get(f'{t}_vagas'))
        f.campo(f'{t}_un', f'{nome_item} {i} — Unidades', 'unidades (inteiro > 0)', f'tipologias[{i}].quantidade',
                F_INT, minimo=0, inteiro=True, ex=ex.get(f'{t}_un'))
        f.campo(f'{t}_preco', f'{nome_item} {i} — Preço/m² sugerido', 'R$/m² (preço padrão da alocação)',
                f'tipologias[{i}].preco_m2', F_RS, minimo=0, ex=ex.get(f'{t}_preco'))

    # ── Grupos de venda ──
    N_GRP = 5
    f.secao(f'6 · Grupos de venda (Viabilidade › Receitas) — {N_GRP} vagas; absorção = % das unidades vendido em cada período')
    f.nota('Pós-chaves = 100% − Pré − Lançamento − Obra e Repasse = 100% − Σ % do grupo na seção de Pagamento: ambos calculados '
           '(zona de referência). Sem pré-lançamento, deixe a absorção de pré em branco.')
    f.nota('Saldo de chaves = o Repasse acima. "Repasse bancário" é a parte dele paga pelo banco no mês seguinte ao fim da obra '
           '(vazio = 100%, tudo no banco). O resto vira carteira própria: a empresa financia o comprador em N parcelas mensais '
           'Price com os Juros de tabela, a partir do mês seguinte ao fim da obra + a carência. Loteamento sem banco: 0%.')
    for g in range(1, N_GRP + 1):
        k = f'gr{g}'
        f.grupo(f'Grupo de venda {g}')
        f.campo(f'{k}_nome', f'Grupo {g} — Nome', 'texto', f'grupos[{g}].nome', F_TXT,
                ex=ex.get(f'{k}_nome'))
        f.campo(f'{k}_pre', f'Grupo {g} — Absorção no Pré-lançamento (só com pré)', '% das unidades',
                f'grupos[{g}].pre_lancamento_pct', F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_pre'))
        f.campo(f'{k}_lanc', f'Grupo {g} — Absorção no Lançamento', '% das unidades', f'grupos[{g}].lancamento_pct',
                F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_lanc'))
        f.campo(f'{k}_obra', f'Grupo {g} — Absorção Durante a obra', '% das unidades', f'grupos[{g}].obra_pct',
                F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_obra'))
        f.campo(f'{k}_janela', f'Grupo {g} — Janela pós-chaves', 'meses (1 a 12)', f'grupos[{g}].pos_chaves_meses', F_INT,
                minimo=1, maximo=12, inteiro=True, ex=ex.get(f'{k}_janela'))
        f.campo(f'{k}_corr', f'Grupo {g} — Correção de estoque', 'Sim/Não', f'grupos[{g}].correcao_estoque', F_TXT, lista=SN,
                ex=ex.get(f'{k}_corr'))
        f.campo(f'{k}_resid', f'Grupo {g} — Resíduo sem prazo', 'lista: imediato | concentrado',
                f'grupos[{g}].residuo_sem_prazo', F_TXT, lista=['imediato', 'concentrado'], ex=ex.get(f'{k}_resid'))
        f.campo(f'{k}_banco', f'Grupo {g} — Repasse bancário', '% do saldo de chaves (vazio = 100%)',
                f'grupos[{g}].repasse_banco_pct', F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_banco'))
        f.campo(f'{k}_cart_parc', f'Grupo {g} — Carteira própria: nº de parcelas (só com repasse < 100%)', 'inteiro (≥ 1)',
                f'grupos[{g}].carteira_parcelas', F_INT, minimo=1, inteiro=True, ex=ex.get(f'{k}_cart_parc'))
        f.campo(f'{k}_cart_car', f'Grupo {g} — Carteira própria: carência após a entrega', 'meses (inteiro ≥ 0; vazio = 0)',
                f'grupos[{g}].carteira_carencia_meses', F_INT, minimo=0, inteiro=True, ex=ex.get(f'{k}_cart_car'))

    # ── Pagamento ──
    N_PAG = 10
    f.secao(f'7 · Fluxo de pagamento por grupo — {N_PAG} vagas (uma linha por condição; até 4 de Parcelamento por grupo)')
    f.nota('Os juros vêm do campo "Juros de tabela". Nome do grupo exatamente como na seção 6. Σ % do grupo ≤ 100%; o que sobrar é o saldo de chaves (banco × carteira própria: seção 6).')
    for j in range(1, N_PAG + 1):
        k = f'pg{j}'
        f.grupo(f'Condição de pagamento {j}')
        f.campo(f'{k}_grupo', f'Condição {j} — Grupo', 'nome do grupo (seção 6)', f'pagamento[{j}].grupo', F_TXT,
                ex=ex.get(f'{k}_grupo'))
        f.campo(f'{k}_linha', f'Condição {j} — Linha', 'lista: Entrada | Parcelamento', f'pagamento[{j}].linha', F_TXT,
                lista=['Entrada', 'Parcelamento'], ex=ex.get(f'{k}_linha'))
        f.campo(f'{k}_pct', f'Condição {j} — % do total', '% do preço', f'pagamento[{j}].pct', F_PCT, minimo=0, maximo=1,
                ex=ex.get(f'{k}_pct'))
        f.campo(f'{k}_parc', f'Condição {j} — Nº de parcelas', 'inteiro (vazio se ao longo da obra)', f'pagamento[{j}].parcelas',
                F_INT, minimo=1, inteiro=True, ex=ex.get(f'{k}_parc'))
        f.campo(f'{k}_desc', f'Condição {j} — Desconto (só Entrada)', '% sobre a entrada', f'pagamento[{j}].desconto_pct',
                F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_desc'))
        f.campo(f'{k}_sinal', f'Condição {j} — Sinal (só Parcelamento)', '% do parcelamento', f'pagamento[{j}].sinal_pct',
                F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_sinal'))
        f.campo(f'{k}_obra', f'Condição {j} — Ao longo da obra (só Parcelamento)', 'Sim/Não', f'pagamento[{j}].ao_longo_obra',
                F_TXT, lista=SN, ex=ex.get(f'{k}_obra'))

    # ── Alocações ──
    N_AL = 10
    f.secao(f'8 · Alocações de venda ({"lote" if lot else "tipologia"} × grupo) — {N_AL} vagas; Σ unidades alocadas por item ≤ catálogo')
    for j in range(1, N_AL + 1):
        k = f'al{j}'
        f.grupo(f'Alocação {j}')
        f.campo(f'{k}_grupo', f'Alocação {j} — Grupo', 'nome do grupo (seção 6)', f'alocacoes[{j}].grupo', F_TXT,
                ex=ex.get(f'{k}_grupo'))
        f.campo(f'{k}_tip', f'Alocação {j} — {nome_item}', f'nome do {"lote" if lot else "tipologia"} (seção 5)',
                f'alocacoes[{j}].tipologia', F_TXT, ex=ex.get(f'{k}_tip'))
        f.campo(f'{k}_un', f'Alocação {j} — Unidades', 'unidades (inteiro)', f'alocacoes[{j}].unidades', F_INT, minimo=0,
                inteiro=True, ex=ex.get(f'{k}_un'))
        f.campo(f'{k}_preco', f'Alocação {j} — Preço', 'R$/m² (sobre privativa + aberta)', f'alocacoes[{j}].preco_m2', F_RS,
                minimo=0, ex=ex.get(f'{k}_preco'))

    # ── Custos ──
    N_CU = 28
    f.secao(f'9 · Custos (Viabilidade › Custos) — {N_CU} vagas + permutas do terreno')
    f.nota(NOTA_CUSTOS)
    f.nota('Em cada vaga: escolha a Unidade e preencha o Valor na linha certa — "Valor (R$ ou R$/m²)" quando a unidade for rs, '
           'rs_m2_priv ou rs_m2_terreno; "Valor (%)" quando for pct_*. A coluna Unidade da linha de valor mostra o que a '
           'unidade escolhida significa. Vagas 1 a 3 já vêm com Grupo e Categoria obrigatórios (cinza).')
    fixos = {1: ('terreno', 'Preço'), 2: ('obra', 'Construção'), 3: ('diretos', 'Corretagem de vendas')}
    for k_ in range(1, N_CU + 1):
        c = f'cu{k_}'
        f.grupo(f'Custo {k_}' + {1: ' — obrigatório: terreno / Preço', 2: ' — obrigatório: obra / Construção',
                                 3: ' — obrigatório: diretos / Corretagem de vendas (% VGV)'}.get(k_, ''))
        fx = fixos.get(k_)
        f.campo(f'{c}_grupo', f'Custo {k_} — Grupo', 'lista: terreno | obra | diretos | indireto | financeiro',
                f'custos[{k_}].grupo', F_TXT,
                lista=['terreno', 'obra', 'diretos', 'indireto', 'financeiro'], fixo=fx[0] if fx else None,
                ex=ex.get(f'{c}_grupo'))
        f.campo(f'{c}_cat', f'Custo {k_} — Categoria', 'ver tabela acima', f'custos[{k_}].categoria', F_TXT,
                lista_ref='categorias', aviso=True, fixo=fx[1] if fx else None, ex=ex.get(f'{c}_cat'))
        f.campo(f'{c}_sub', f'Custo {k_} — Subcategoria / descrição', 'Preço: Valor à vista | Parcelado; Outro: texto',
                f'custos[{k_}].subcategoria', F_TXT, ex=ex.get(f'{c}_sub'))
        f.campo(f'{c}_un', f'Custo {k_} — Unidade', 'lista: ' + ' | '.join(UNIDADES_CUSTO),
                f'custos[{k_}].unidade', F_TXT, lista=UNIDADES_CUSTO,
                fixo='pct_vgv' if k_ == 3 else None, ex=ex.get(f'{c}_un'))
        f.campo(f'{c}_val', f'Custo {k_} — Valor (R$ ou R$/m²)', ROTULO_UNIDADE.replace('{u}', '{' + f'{c}_un' + '}'),
                f'custos[{k_}].valor_rs', F_RS, minimo=0, ex=ex.get(f'{c}_val'))
        f.campo(f'{c}_pct', f'Custo {k_} — Valor (%)', ROTULO_UNIDADE_PCT.replace('{u}', '{' + f'{c}_un' + '}'),
                f'custos[{k_}].valor_pct', F_PCT, minimo=0, maximo=10,
                ex=ex.get(f'{c}_pct'))
        f.campo(f'{c}_dist', f'Custo {k_} — Distribuição', 'Linear | nome de uma curva do catálogo (ex.: Curva S) | unit_delivery | sales_revenue',
                f'custos[{k_}].distribuicao', F_TXT, lista_ref='distribuicao', aviso=True, ex=ex.get(f'{c}_dist'))
        f.campo(f'{c}_cron', f'Custo {k_} — Cronograma', 'evento, customizado ou "fase: <nome>"',
                f'custos[{k_}].cronograma', F_TXT, lista_ref='cronograma', aviso=True,
                ex=ex.get(f'{c}_cron'))
        f.campo(f'{c}_ini', f'Custo {k_} — Início (só customizado)', 'mês (inteiro ≥ 0)',
                f'custos[{k_}].inicio_mes', F_INT, minimo=0, inteiro=True, ex=ex.get(f'{c}_ini'))
        f.campo(f'{c}_dur', f'Custo {k_} — Duração (só customizado)', 'meses (inteiro ≥ 1)',
                f'custos[{k_}].duracao_meses', F_INT, minimo=1, inteiro=True, ex=ex.get(f'{c}_dur'))

    f.grupo('Permuta física do terreno (terreno / Preço / Permuta física) — até 2')
    for m in (1, 2):
        pf = f'pf{m}'
        f.campo(f'{pf}_tip', f'Permuta física {m} — {nome_item} entregue', 'nome do ' + ('lote' if lot else 'tipologia'),
                f'permuta_fisica[{m}].tipologia', F_TXT, ex=ex.get(f'{pf}_tip'))
        f.campo(f'{pf}_un', f'Permuta física {m} — Unidades entregues', 'unidades (≤ alocado)',
                f'permuta_fisica[{m}].unidades', F_INT, minimo=0, inteiro=True,
                ex=ex.get(f'{pf}_un'))
    f.grupo('Permuta financeira do terreno (terreno / Preço / Permuta financeira) — até 2')
    for m in (1, 2):
        pn = f'pn{m}'
        f.campo(f'{pn}_un', f'Permuta financeira {m} — Unidade', 'lista: rs | pct_vgv',
                f'permuta_financeira[{m}].unidade', F_TXT, lista=['rs', 'pct_vgv'],
                ex=ex.get(f'{pn}_un'))
        f.campo(f'{pn}_val', f'Permuta financeira {m} — Valor em R$', 'R$ (valor cheio; se unidade = rs)',
                f'permuta_financeira[{m}].valor_rs', F_RS, minimo=0, ex=ex.get(f'{pn}_val'))
        f.campo(f'{pn}_pct', f'Permuta financeira {m} — Valor em % do VGV', '% do VGV (se unidade = pct_vgv)',
                f'permuta_financeira[{m}].valor_pct', F_PCT, minimo=0, maximo=1,
                ex=ex.get(f'{pn}_pct'))
        f.campo(f'{pn}_imp', f'Permuta financeira {m} — Deduzir impostos', 'Sim/Não',
                f'permuta_financeira[{m}].deduzir_imposto', F_TXT, lista=SN,
                ex=ex.get(f'{pn}_imp'))
        f.campo(f'{pn}_cor', f'Permuta financeira {m} — Deduzir corretagem', 'Sim/Não',
                f'permuta_financeira[{m}].deduzir_corretagem', F_TXT, lista=SN,
                ex=ex.get(f'{pn}_cor'))

    # ── Funding ──
    f.secao('10 · Funding (operações) — 1 financiamento à produção + 3 dívidas + 3 equities + 3 antecipações de recebíveis; '
            'deixe em branco o que não existir')
    f.grupo('Financiamento à produção (no máximo um por estudo)')
    f.campo('fp_nome', 'Financiamento à produção — Nome', 'texto', 'funding.financiamento_producao.nome', F_TXT,
            ex=ex.get('fp_nome'))
    f.campo('fp_ativo', 'Financiamento à produção — Ativo', 'Sim/Não', 'funding.financiamento_producao.ativo', F_TXT,
            lista=SN, ex=ex.get('fp_ativo'))
    f.campo('fp_valor', 'Financiamento à produção — Teto', 'R$ (0 = sem teto)', 'funding.financiamento_producao.valor',
            F_RS, minimo=0, ex=ex.get('fp_valor'))
    f.campo('fp_taxa', 'Financiamento à produção — Taxa', '% a.a.', 'funding.financiamento_producao.taxa_anual', F_PCT,
            minimo=0, maximo=1, ex=ex.get('fp_taxa'))
    f.campo('fp_exp', 'Financiamento à produção — Exposição mínima', '% (do custo)',
            'funding.financiamento_producao.exposicao_minima', F_PCT, minimo=0, maximo=1, ex=ex.get('fp_exp'))
    f.campo('fp_fin', 'Financiamento à produção — % do custo financiado', '%',
            'funding.financiamento_producao.percentual_financiavel', F_PCT, minimo=0, maximo=1, ex=ex.get('fp_fin'))
    f.campo('fp_caixa', 'Financiamento à produção — Caixa amortiza antes das chaves', 'Sim/Não',
            'funding.financiamento_producao.amortizar_com_caixa_disponivel', F_TXT, lista=SN, ex=ex.get('fp_caixa'))
    f.campo('fp_base', 'Financiamento à produção — Base financiável', 'linhas de custo separadas por ";" — Categoria ou Categoria / Subcategoria; vazio = padrão',
            'funding.financiamento_producao.base_financiavel', F_TXT, ex=ex.get('fp_base'))
    for d in (1, 2, 3):
        k = f'dv{d}'
        f.grupo(f'Dívida {d}')
        f.campo(f'{k}_nome', f'Dívida {d} — Nome', 'texto', f'funding.divida[{d}].nome', F_TXT, ex=ex.get(f'{k}_nome'))
        f.campo(f'{k}_valor', f'Dívida {d} — Valor', 'R$ (valor cheio)', f'funding.divida[{d}].valor', F_RS, minimo=0,
                ex=ex.get(f'{k}_valor'))
        f.campo(f'{k}_aporte', f'Dívida {d} — Mês do aporte', 'evento, fase ("fase: <nome>") ou customizado',
                f'funding.divida[{d}].aporte', F_TXT, lista_ref='cronograma', aviso=True,
                ex=ex.get(f'{k}_aporte'))
        f.campo(f'{k}_mes', f'Dívida {d} — Mês (só customizado)', 'mês (inteiro ≥ 0)', f'funding.divida[{d}].mes',
                F_INT, minimo=0, inteiro=True, ex=ex.get(f'{k}_mes'))
        f.campo(f'{k}_dist', f'Dívida {d} — Distribuir aporte', 'Sim/Não', f'funding.divida[{d}].distribuir_aporte',
                F_TXT, lista=SN, ex=ex.get(f'{k}_dist'))
        f.campo(f'{k}_aporte_m', f'Dívida {d} — Aporte em', 'meses (≥ 1)', f'funding.divida[{d}].aporte_meses', F_INT,
                minimo=1, inteiro=True, ex=ex.get(f'{k}_aporte_m'))
        f.campo(f'{k}_taxa', f'Dívida {d} — Taxa', '% a.a.', f'funding.divida[{d}].taxa_anual', F_PCT, minimo=0,
                maximo=1, ex=ex.get(f'{k}_taxa'))
        f.campo(f'{k}_amort', f'Dívida {d} — Amortização', 'meses', f'funding.divida[{d}].amortizacao_meses',
                F_INT, minimo=0, inteiro=True, ex=ex.get(f'{k}_amort'))
        f.campo(f'{k}_carencia', f'Dívida {d} — Carência', 'meses (< amortização)',
                f'funding.divida[{d}].carencia_meses', F_INT, minimo=0, inteiro=True, ex=ex.get(f'{k}_carencia'))
        f.campo(f'{k}_estr', f'Dívida {d} — Estruturação', '% do valor', f'funding.divida[{d}].estruturacao_pct',
                F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_estr'))
        f.campo(f'{k}_adm', f'Dívida {d} — Administração', 'R$/mês', f'funding.divida[{d}].administracao_mensal',
                F_RS, minimo=0, ex=ex.get(f'{k}_adm'))
        f.campo(f'{k}_outros', f'Dívida {d} — Outros encargos', 'R$ (iniciais)', f'funding.divida[{d}].outros_encargos',
                F_RS, minimo=0, ex=ex.get(f'{k}_outros'))
    for q in (1, 2, 3):
        k = f'eq{q}'
        f.grupo(f'Equity {q}')
        f.campo(f'{k}_nome', f'Equity {q} — Nome', 'texto', f'funding.equity[{q}].nome', F_TXT, ex=ex.get(f'{k}_nome'))
        f.campo(f'{k}_valor', f'Equity {q} — Valor', 'R$ (valor cheio)', f'funding.equity[{q}].valor', F_RS, minimo=0,
                ex=ex.get(f'{k}_valor'))
        f.campo(f'{k}_aporte', f'Equity {q} — Mês do aporte', 'evento, fase ("fase: <nome>") ou customizado',
                f'funding.equity[{q}].aporte', F_TXT, lista_ref='cronograma', aviso=True,
                ex=ex.get(f'{k}_aporte'))
        f.campo(f'{k}_mes', f'Equity {q} — Mês (só customizado)', 'mês (inteiro ≥ 0)', f'funding.equity[{q}].mes',
                F_INT, minimo=0, inteiro=True, ex=ex.get(f'{k}_mes'))
        f.campo(f'{k}_modo', f'Equity {q} — Modo de retorno', 'lista: permuta_financeira | resultado_final',
                f'funding.equity[{q}].modo_retorno', F_TXT, lista=['permuta_financeira', 'resultado_final'],
                ex=ex.get(f'{k}_modo'))
        f.campo(f'{k}_pct', f'Equity {q} — % de retorno', '% (do resultado final ou da receita líquida)',
                f'funding.equity[{q}].pct_retorno', F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_pct'))

    f.nota('Antecipação de recebíveis: cessão de parte da carteira a securitizadora/CRI/FIDC. Não tem campo de valor — o app '
           'calcula o valor presente das parcelas cedidas à taxa de desconto. Com âncora num evento, a cessão é no INÍCIO dele; '
           'cedem-se as parcelas que vencem depois desse mês, das vendas feitas até ele. Uma antecipação posterior cede o % do '
           'que a anterior não cedeu. Entra só no Funding/Fluxo alavancado, nunca na Proforma nem na TIR/VPL do projeto.')
    for a_ in (1, 2, 3):
        k = f'an{a_}'
        f.grupo(f'Antecipação de recebíveis {a_}')
        f.campo(f'{k}_nome', f'Antecipação {a_} — Nome', 'texto', f'funding.antecipacao[{a_}].nome', F_TXT, ex=ex.get(f'{k}_nome'))
        f.campo(f'{k}_aporte', f'Antecipação {a_} — Mês da cessão', 'evento, fase ("fase: <nome>") ou customizado',
                f'funding.antecipacao[{a_}].aporte', F_TXT, lista_ref='cronograma', aviso=True, ex=ex.get(f'{k}_aporte'))
        f.campo(f'{k}_mes', f'Antecipação {a_} — Mês (só customizado)', 'mês (inteiro ≥ 0)', f'funding.antecipacao[{a_}].mes',
                F_INT, minimo=0, inteiro=True, ex=ex.get(f'{k}_mes'))
        f.campo(f'{k}_pct', f'Antecipação {a_} — Carteira cedida', '% do que as vendas até a cessão ainda vão pagar',
                f'funding.antecipacao[{a_}].pct_cedida', F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_pct'))
        f.campo(f'{k}_taxa', f'Antecipação {a_} — Taxa de desconto', '% a.a.', f'funding.antecipacao[{a_}].taxa_anual', F_PCT,
                minimo=0, maximo=1, ex=ex.get(f'{k}_taxa'))
        f.campo(f'{k}_estr', f'Antecipação {a_} — Estruturação', '% do valor presente (no mês da cessão)',
                f'funding.antecipacao[{a_}].estruturacao_pct', F_PCT, minimo=0, maximo=1, ex=ex.get(f'{k}_estr'))
        f.campo(f'{k}_adm', f'Antecipação {a_} — Administração', 'R$/mês (enquanto houver parcela cedida)',
                f'funding.antecipacao[{a_}].administracao_mensal', F_RS, minimo=0, ex=ex.get(f'{k}_adm'))
        f.campo(f'{k}_outros', f'Antecipação {a_} — Outros custos', 'R$ (no mês da cessão)',
                f'funding.antecipacao[{a_}].outros_encargos', F_RS, minimo=0, ex=ex.get(f'{k}_outros'))

    # ── Cenários ──
    N_CEN = 5
    f.secao(f'11 · Cenários (variação sobre o estudo real) — {N_CEN} vagas')
    for c_ in range(1, N_CEN + 1):
        k = f'ce{c_}'
        f.campo(f'{k}_nome', f'Cenário {c_} — Nome', 'texto', f'cenarios[{c_}].nome', F_TXT, ex=ex.get(f'{k}_nome'))
        f.campo(f'{k}_preco', f'Cenário {c_} — Preço de venda', 'variação % (pode ser negativa)',
                f'cenarios[{c_}].preco_venda_pct', F_PCT, minimo=-1, maximo=10, ex=ex.get(f'{k}_preco'))
        f.campo(f'{k}_obra', f'Cenário {c_} — Custo de obra', 'variação % (pode ser negativa)',
                f'cenarios[{c_}].custo_obra_pct', F_PCT, minimo=-1, maximo=10, ex=ex.get(f'{k}_obra'))

    # ════════ referência calculada ════════
    f.secao('REFERÊNCIA CALCULADA — não preencher (conferência do que foi preenchido)', calc=True)
    f.nota('TIR, VPL, payback, Proforma e Fluxo de Caixa NÃO são calculados aqui: dependem do motor mensal do app. '
           'Abaixo, só o que se deduz das entradas: cronograma, áreas, VGV de tabela, absorção, repasse e custos convertidos em R$.')

    f.grupo('Cronograma — inícios travados e datas')
    f.calc('k_plan_fim', 'Planejamento — fim', 'mês', f'N({r("plan_ini")})+N({r("plan_dur")})', F_INT)
    f.calc('k_pre_ini', 'Pré-lançamento — Início (travado = fim do planejamento)', 'mês',
           f'IF({r("tem_pre")}="Sim",{r("k_plan_fim")},"—")', F_INT, chave='cronograma.pre_lancamento.inicio_mes')
    f.calc('k_pre_fim', 'Pré-lançamento — fim', 'mês', f'IF({r("tem_pre")}="Sim",{r("k_pre_ini")}+N({r("pre_dur")}),"—")', F_INT)
    f.calc('k_lanc_ini', 'Lançamento — Início (travado = fim do evento anterior)', 'mês',
           f'IF({r("tem_pre")}="Sim",{r("k_pre_fim")},{r("k_plan_fim")})', F_INT, chave='cronograma.lancamento.inicio_mes')
    f.calc('k_lanc_fim', 'Lançamento — fim', 'mês', f'{r("k_lanc_ini")}+N({r("lanc_dur")})', F_INT)
    f.calc('k_obra_fim', 'Obra — fim', 'mês', f'N({r("obra_ini")})+N({r("obra_dur")})', F_INT)
    f.calc('k_pos_ini', 'Pós-obras — Início (travado = fim da obra)', 'mês', r('k_obra_fim'), F_INT,
           chave='cronograma.pos_obra.inicio_mes')
    f.calc('k_pos_fim', 'Pós-obras — fim = fim do projeto', 'mês', f'{r("k_pos_ini")}+N({r("pos_dur")})', F_INT)
    for nome, rot, mes in (('k_d_lanc', 'Data do Lançamento', 'k_lanc_ini'), ('k_d_obra', 'Data de início da Obra', 'obra_ini'),
                           ('k_d_chaves', 'Data das chaves (fim da Obra)', 'k_obra_fim'), ('k_d_fim', 'Data de fim do projeto', 'k_pos_fim')):
        f.calc(nome, rot, 'mês/ano', f'IF(ISNUMBER({r("data_ini")}),EDATE({r("data_ini")},{r(mes)}),"")', F_DATA)
    f.calc('k_a_obra', 'Obra começa antes do fim do Lançamento?', 'alerta',
           f'IF(N({r("obra_ini")})<{r("k_lanc_fim")},"⚠ obra inicia antes do fim do lançamento","OK")')

    itens = range(1, N_TIP + 1)
    f.grupo(f'{"Lotes" if lot else "Tipologias"} — catálogo')
    for i in itens:
        t = f'tip{i}'
        f.calc(f'{t}_cat', f'{nome_item} {i} — VGV de catálogo (preço sugerido)', 'R$',
               f'(N({r(t+"_priv")})+N({r(t+"_aberta")}))*N({r(t+"_un")})*N({r(t+"_preco")})', F_RS)
    f.calc('k_un_cat', 'Unidades no catálogo', 'unidades', '+'.join(f'N({r(f"tip{i}_un")})' for i in itens), F_INT)
    f.calc('k_vgv_cat', 'VGV de catálogo (preço sugerido)', 'R$', '+'.join(r(f'tip{i}_cat') for i in itens), F_RS)

    f.grupo('Alocações — VGV de tabela por alocação (unidades × (privativa+aberta) × preço)')
    for j in range(1, N_AL + 1):
        a = f'al{j}'
        area_item = '+'.join(f'({r(f"tip{i}_nome")}={r(a+"_tip")})*(N({r(f"tip{i}_priv")})+N({r(f"tip{i}_aberta")}))' for i in itens)
        f.calc(f'{a}_area', f'Alocação {j} — área (privativa+aberta) × unidades', 'm²',
               f'IF({r(a+"_tip")}="",0,({area_item})*N({r(a+"_un")}))', F_M2)
        f.calc(f'{a}_vgv', f'Alocação {j} — VGV de tabela', 'R$', f'{r(a+"_area")}*N({r(a+"_preco")})', F_RS)
    f.calc('k_area_aloc', 'Área privativa total alocada (base de R$/m² priv.)', 'm²',
           '+'.join(r(f'al{j}_area') for j in range(1, N_AL + 1)), F_M2)
    f.calc('k_un_aloc', 'Unidades alocadas', 'unidades', '+'.join(f'N({r(f"al{j}_un")})' for j in range(1, N_AL + 1)), F_INT)
    f.calc('k_vgv', 'VGV de tabela (Σ alocações)', 'R$', '+'.join(r(f'al{j}_vgv') for j in range(1, N_AL + 1)), F_RS)
    f.calc('k_pm2', 'Preço médio de tabela', 'R$/m² priv.', f'IF({r("k_area_aloc")}>0,{r("k_vgv")}/{r("k_area_aloc")},"")', F_RS)
    for i in itens:
        t = f'tip{i}'
        f.calc(f'{t}_aloc', f'{nome_item} {i} — unidades alocadas', 'unidades',
               '+'.join(f'({r(f"al{j}_tip")}={r(t+"_nome")})*N({r(f"al{j}_un")})' for j in range(1, N_AL + 1)), F_INT)
        f.calc(f'{t}_alerta', f'{nome_item} {i} — alocado × catálogo', 'alerta',
               f'IF({r(t+"_nome")}="","",IF({r(t+"_aloc")}>N({r(t+"_un")}),"⚠ alocado > catálogo",'
               f'IF({r(t+"_aloc")}<N({r(t+"_un")}),"há unidades sem alocar","OK")))')
    for j in range(1, N_AL + 1):
        a = f'al{j}'
        achou = '+'.join(f'({r(f"tip{i}_nome")}={r(a+"_tip")})*1' for i in itens)
        f.calc(f'{a}_alerta', f'Alocação {j} — item encontrado?', 'alerta',
               f'IF({r(a+"_tip")}="","",IF(({achou})>0,"OK","⚠ nome não bate com a seção 5"))')

    f.grupo('Grupos de venda — absorção, repasse e VGV')
    for g in range(1, N_GRP + 1):
        k = f'gr{g}'
        soma_pg = '+'.join(f'IF({r(f"pg{j}_grupo")}={r(k+"_nome")},N({r(f"pg{j}_pct")}),0)' for j in range(1, N_PAG + 1))
        f.calc(f'{k}_pos', f'Grupo {g} — Absorção Pós-chaves', '% (= 100% − pré − lançamento − obra)',
               f'IF({r(k+"_nome")}="","",1-N({r(k+"_pre")})-N({r(k+"_lanc")})-N({r(k+"_obra")}))', F_PCT, chave='calculado (absorcao pós-chaves)')
        f.calc(f'{k}_sp', f'Grupo {g} — Σ % do pagamento', '% do preço', f'IF({r(k+"_nome")}="","",{soma_pg})', F_PCT)
        f.calc(f'{k}_rep', f'Grupo {g} — Repasse', '% (= 100% − Σ pagamento)', f'IF({r(k+"_nome")}="","",1-{r(k+"_sp")})', F_PCT,
               chave='calculado (repasse)')
        f.calc(f'{k}_cart', f'Grupo {g} — Carteira própria', '% do preço (= Repasse × (100% − Repasse bancário))',
               f'IF({r(k+"_nome")}="","",{r(k+"_rep")}*(1-IF({r(k+"_banco")}="",1,N({r(k+"_banco")}))))', F_PCT,
               chave='calculado (carteira própria)')
        f.calc(f'{k}_vgv', f'Grupo {g} — VGV de tabela', 'R$',
               f'IF({r(k+"_nome")}="","",' + '+'.join(f'IF({r(f"al{j}_grupo")}={r(k+"_nome")},{r(f"al{j}_vgv")},0)' for j in range(1, N_AL + 1)) + ')', F_RS)
        f.calc(f'{k}_alerta', f'Grupo {g} — conferência', 'alerta',
               f'IF({r(k+"_nome")}="","",IF(N({r(k+"_pre")})+N({r(k+"_lanc")})+N({r(k+"_obra")})>1.00001,"⚠ absorção > 100%",'
               f'IF({r(k+"_sp")}>1.00001,"⚠ Σ pagamento > 100%",IF(AND({r("tem_pre")}="Não",N({r(k+"_pre")})>0),"⚠ absorção de pré sem pré-lançamento",'
               f'IF(' + '+'.join(f'IF(AND({r(f"pg{j}_grupo")}={r(k+"_nome")},{r(f"pg{j}_linha")}="Parcelamento"),1,0)' for j in range(1, N_PAG + 1)) +
               '>4,"⚠ mais de 4 parcelamentos",'
               f'IF(AND({r(k+"_banco")}<>"",N({r(k+"_banco")})<0.99999,{r(k+"_cart_parc")}=""),"⚠ repasse < 100% sem nº de parcelas da carteira","OK")))))) ')

    f.grupo('Custos convertidos em R$ (quando dá para converter sem o fluxo)')
    for k_ in range(1, N_CU + 1):
        c = f'cu{k_}'
        f.calc(f'{c}_rs', f'Custo {k_} — equivalente em R$', 'R$ (ou texto se depender do fluxo)',
               f'IF({r(c+"_un")}="","",IF({r(c+"_un")}="rs",N({r(c+"_val")}),IF({r(c+"_un")}="rs_m2_priv",N({r(c+"_val")})*{r("k_area_aloc")},'
               f'IF({r(c+"_un")}="rs_m2_terreno",N({r(c+"_val")})*N({r("terreno_area")}),IF({r(c+"_un")}="pct_vgv",N({r(c+"_pct")})*{r("k_vgv")},'
               f'"depende do fluxo (app)")))))', F_RS)
    # conferência por vaga: categoria × unidade (tabela da tela de Custos), uma só linha de valor, valor na linha certa
    for k_ in range(1, N_CU + 1):
        c = f'cu{k_}'
        g_, ct, un, vr, vp = (r(c + x) for x in ('_grupo', '_cat', '_un', '_val', '_pct'))
        combos = ','.join(
            f'AND({g_}="{grp}",{ct}="{cat}",OR(' + ','.join(f'{un}="{u}"' for u in unidades) + '))'
            for (grp, cat), unidades in UNIDADES_POR_CATEGORIA.items())
        f.calc(f'{c}_alerta', f'Custo {k_} — conferência', 'alerta',
               f'IF(AND({g_}="",{ct}=""),"",IF(OR({g_}="",{ct}="",{un}=""),"⚠ falta grupo, categoria ou unidade",'
               f'IF(NOT(OR({combos})),"⚠ unidade não aceita nesta categoria",'
               f'IF(AND({vr}<>"",{vp}<>""),"⚠ preencha só uma linha de valor",'
               f'IF(IF(LEFT({un},3)="pct",{vp}="",{vr}=""),"⚠ falta o valor na linha certa","OK")))))')
    f.calc('k_1pct', 'Cada 1% do VGV de tabela equivale a (para converter Projetos, que não aceita % do VGV)', 'R$',
           f'{r("k_vgv")}*0.01', F_RS)
    for grp in ('terreno', 'obra', 'diretos', 'indireto', 'financeiro'):
        f.calc(f'k_g_{grp}', f'Σ custos conversíveis — {grp}', 'R$',
               '+'.join(f'IF({r(f"cu{k_}_grupo")}="{grp}",N({r(f"cu{k_}_rs")}),0)' for k_ in range(1, N_CU + 1)), F_RS)
    for m in (1, 2):
        f.calc(f'pf{m}_vgv', f'Permuta física {m} — VGV entregue (unidades × área × preço sugerido)', 'R$',
               f'IF({r(f"pf{m}_tip")}="",0,N({r(f"pf{m}_un")})*(' +
               '+'.join(f'({r(f"tip{i}_nome")}={r(f"pf{m}_tip")})*(N({r(f"tip{i}_priv")})+N({r(f"tip{i}_aberta")}))*N({r(f"tip{i}_preco")})' for i in itens) +
               '))', F_RS)
        f.calc(f'pn{m}_rs', f'Permuta financeira {m} — equivalente em R$', 'R$',
               f'IF({r(f"pn{m}_un")}="pct_vgv",N({r(f"pn{m}_pct")})*{r("k_vgv")},IF({r(f"pn{m}_un")}="rs",N({r(f"pn{m}_val")}),0))', F_RS)
    f.calc('k_dividas', 'Σ valor das dívidas', 'R$', '+'.join(f'N({r(f"dv{d}_valor")})' for d in (1, 2, 3)), F_RS)
    f.calc('k_equity', 'Σ valor dos equities', 'R$', '+'.join(f'N({r(f"eq{q}_valor")})' for q in (1, 2, 3)), F_RS)
    return f


def main(argv):
    saida = Path(argv[1]) if len(argv) > 1 else Path(__file__).resolve().parents[2] / 'referencia' / 'formularios'
    saida.mkdir(parents=True, exist_ok=True)
    for nivel, fabrica in (('preliminar', preliminar), ('avancado', avancado)):
        for tipo in ('incorporacao', 'loteamento'):
            destino = saida / f'formulario-{nivel}-{tipo}.xlsx'
            escrever(fabrica(tipo, {}), destino)
            print(destino)


if __name__ == '__main__':
    main(sys.argv)
