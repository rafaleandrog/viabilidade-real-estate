#!/usr/bin/env python3
"""Cria um estudo na instância a partir de um formulário simplificado preenchido, e confere o resultado.

    python3 -I scripts/formularios/importar_formulario.py <formulario.xlsx> [--nome-estudo "qa-..."]
                                                         [--executar] [--relatorio saida.json]

Sem --executar valida a planilha e mostra o plano (nada é escrito). Com --executar, nesta ordem:

    1. VALIDA a planilha (fail-closed: erro de preenchimento para aqui, antes de escrever qualquer coisa);
    2. CRIA o estudo pelas rotas do app (URBIVERSO_QA_URL + URBIVERSO_QA_TOKEN_PCPAL; o token nunca é impresso);
    3. LEVA o pagamento ao contrato canônico (`aplicar-pagamento-canonico.ts`) — necessário enquanto a instância
       rodar um backend anterior à normalização na API; depois dela é inofensivo;
    4. RELÊ o que foi gravado e compara com o que a planilha pediu;
    5. RODA o motor do repositório sobre o estudo (`conferir-estudo.ts`) e separa o que é defeito do estudo do que
       é achado conhecido do próprio app.

Sai com código 0 só se não houver erro de validação, divergência na releitura nem erro de invariante do app.

O que é lido da planilha: coluna A (rótulo), B (valor), C (unidade), D (chave técnica) e o FORMATO da célula B
(célula formatada em % guarda fração: 0,04 vira 4 na API). Nada vem do gerador. A chave técnica é a única fonte
de decisão: `estudos.<coluna>`, `parametros.<c>`, `cronograma.<evento>.<campo>`, `<coluna>_modo = <modo> → <coluna>_valor`
e `<lista>[n].<campo>` (`funding.divida[2].taxa_anual`). Chave sem tratamento vai para a lista de IGNORADOS, nunca some calada.

Requisitos: python3, openpyxl, e — para os passos 3 e 5 — node com as dependências do repositório (`pnpm install`).
"""
import argparse
import http.client
import json
import os
import re
import ssl
import subprocess
import sys
import urllib.error
import urllib.request
from collections import defaultdict
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
CTX = ssl.create_default_context(cafile='/root/.ccr/ca-bundle.crt') if Path('/root/.ccr/ca-bundle.crt').exists() else ssl.create_default_context()

ALT = re.compile(r'^(\w+) = (\w+) → (\w+)')
LISTA = re.compile(r'^([a-z_.]+)\[(\d+)\]\.([a-z0-9_]+)$')
MESES = ['jan', 'fev', 'mar', 'abr', 'mai', 'jun', 'jul', 'ago', 'set', 'out', 'nov', 'dez']

# chaves que o importador conhece e NÃO envia, com o motivo (aparecem no relatório quando preenchidas)
NAO_ENVIADAS = {
    'estudos.regiao_mercado_id': 'região de mercado exige o id da região cadastrada; escolha pela tela',
    'apelo_comercial_documentos.tipo_dado': 'apelo comercial é documento anexado; envie pela tela',
    'apelo_comercial_documentos.texto_adicional': 'apelo comercial é documento anexado; envie pela tela',
    'estudo_documentos': 'a imagem principal é upload; envie pela tela',
}

UNIDADES_POR_CATEGORIA = {
    ('terreno', 'Preço'): ['rs', 'rs_m2_terreno'], ('terreno', 'Registro'): ['rs', 'rs_m2_priv'], ('terreno', 'Outro'): ['rs', 'pct_vgv'],
    ('obra', 'Construção'): ['rs', 'rs_m2_priv'], ('obra', 'Outorga'): ['rs', 'pct_vgv'], ('obra', 'Decoração'): ['rs', 'rs_m2_priv'],
    ('obra', 'Gestão da obra'): ['rs', 'pct_obra'], ('obra', 'Contingência'): ['rs', 'pct_vgv'], ('obra', 'Outro'): ['rs', 'pct_vgv'],
    ('diretos', 'Marketing & Publicidade'): ['rs', 'pct_vgv', 'pct_recebido'], ('diretos', 'Corretagem de vendas'): ['pct_vgv'],
    ('diretos', 'Projetos'): ['rs', 'rs_m2_priv', 'pct_constr'], ('diretos', 'Licenças e Aprovações'): ['rs', 'rs_m2_priv'],
    ('diretos', 'Manutenção pós-obra'): ['rs', 'pct_vgv', 'pct_recebido'], ('diretos', 'Outro'): ['rs', 'pct_vgv', 'pct_recebido'],
    ('indireto', 'Marketing global'): ['rs', 'pct_vgv', 'pct_recebido'], ('indireto', 'Stand de vendas'): ['rs', 'pct_vgv'],
    ('indireto', 'Gestão'): ['rs', 'pct_vgv', 'pct_recebido'], ('indireto', 'Outro'): ['rs', 'pct_vgv', 'pct_recebido'],
    ('financeiro', 'Juros de financiamento'): ['rs', 'pct_vgv'], ('financeiro', 'Taxas bancárias'): ['rs', 'pct_vgv'],
    ('financeiro', 'Estruturação de dívida'): ['rs', 'pct_vgv'], ('financeiro', 'Investidores'): ['rs', 'pct_vgv'],
    ('financeiro', 'Outro'): ['rs', 'pct_vgv'],
}
EVENTOS = ['planejamento', 'pre_lancamento', 'lancamento', 'obra', 'pos_obra']

# achados do próprio app, já registrados — não reprovam o estudo, mas aparecem
CONHECIDOS = [
    (re.compile(r'VGV Σ (tipologias|alocações)'), '#835 (conferir-estudo ignora área aberta)'),
    (re.compile(r'Σ legado entrada\+parcelas'), '#827 (repasse derivado)'),
    (re.compile(r'CAMADAS_DIVERGEM_PERMUTA_FISICA'), '#832 (camada Premissas inexistente no Avançado)'),
    (re.compile(r'RECEITA_BRUTA_NAO_CONSERVA'), '#824 (resíduo de centavos)'),
    (re.compile(r'pós-chaves .*GRAVADO vs DERIVADO'), '#840 (default de absorção)'),
]


# ═════════════════════════ leitura ═════════════════════════
class Entrada:
    __slots__ = ('rotulo', 'valor', 'chave', 'pct')

    def __init__(self, rotulo, valor, chave, pct):
        self.rotulo, self.valor, self.chave, self.pct = rotulo, valor, chave, pct


def ler(caminho):
    wb = openpyxl.load_workbook(caminho)
    if 'Formulário' not in wb.sheetnames:
        raise SystemExit(f'{caminho}: não tem a aba "Formulário" (abas: {", ".join(wb.sheetnames)}). Use um dos modelos de referencia/formularios/.')
    ws = wb['Formulário']
    entradas = []
    for row in ws.iter_rows(min_row=4):
        a, b, d = row[0], row[1], row[3]
        if a.value is None:
            continue
        if str(a.value).startswith('REFERÊNCIA CALCULADA'):
            break
        if d.value is None:
            continue
        v = b.value
        if v is None or v == '':
            continue
        pct = '%' in (b.number_format or '')
        if isinstance(v, (int, float)) and not isinstance(v, bool) and pct:
            v = round(v * 100, 8)
        elif v in ('Sim', 'Não'):
            v = v == 'Sim'
        elif hasattr(v, 'strftime'):
            v = f'{MESES[v.month - 1]}/{v.year}'
        entradas.append(Entrada(str(a.value), v, str(d.value), pct))
    return entradas


class Modelo:
    def __init__(self):
        self.estudo, self.param, self.crono = {}, {}, {}
        self.listas = defaultdict(lambda: defaultdict(dict))
        self.alternativas = defaultdict(list)      # coluna_modo -> [rótulo, ...]
        self.nao_enviadas = []                     # (rótulo, motivo)
        self.ignoradas = []                        # (rótulo, chave)
        self.rotulos = {}                          # chave -> rótulo, para mensagens

    def __getitem__(self, lista):
        return self.listas[lista]


def montar(entradas):
    m = Modelo()
    for e in entradas:
        m.rotulos[e.chave] = e.rotulo
        k = e.chave
        a = ALT.match(k)
        if a:
            modo_col, modo, valor_col = a.groups()
            m.estudo[modo_col] = modo
            m.estudo[valor_col] = e.valor
            m.alternativas[modo_col].append(e.rotulo)
        elif k in NAO_ENVIADAS:
            m.nao_enviadas.append((e.rotulo, NAO_ENVIADAS[k]))
        elif re.fullmatch(r'estudos\.\w+', k):
            m.estudo[k.split('.', 1)[1]] = e.valor
        elif k.startswith('parametros.'):
            m.param[k.split('.', 1)[1]] = e.valor
        elif re.fullmatch(r'cronograma\.\w+\.\w+', k):
            _, ev, campo = k.split('.')
            m.crono.setdefault(ev, {})[campo] = e.valor
        elif LISTA.match(k):
            lista, n, campo = LISTA.match(k).groups()
            m.listas[lista][int(n)][campo] = e.valor
        elif k.startswith('funding.financiamento_producao.'):
            m.listas['funding.financiamento_producao'][1][k.rsplit('.', 1)[1]] = e.valor
        elif k in NAO_ENVIADAS or k.split(' ')[0] in NAO_ENVIADAS:
            m.nao_enviadas.append((e.rotulo, NAO_ENVIADAS.get(k) or NAO_ENVIADAS[k.split(' ')[0]]))
        elif k.startswith('estudo_imoveis.'):
            m.nao_enviadas.append((e.rotulo, 'vínculo com imóvel do Núcleo não é suportado; use origem = manual'))
        else:
            m.ignoradas.append((e.rotulo, k))
    return m


# ═════════════════════════ validação (fail-closed) ═════════════════════════
def rotulos_da_base(texto):
    """'Preço / Valor à vista; Construção' -> [('Preço', 'Valor à vista'), ('Construção', None)]."""
    saida = []
    for parte in [x.strip() for x in str(texto or '').split(';') if x.strip()]:
        cat, _, sub = parte.partition('/')
        saida.append((cat.strip(), sub.strip() or None))
    return saida


def linhas_de_custo_planejadas(m):
    """(categoria, subcategoria) de cada linha de custo que a planilha vai criar, na ordem das vagas."""
    linhas = [(c.get('categoria'), c.get('subcategoria') or None) for _, c in ordenados(m['custos']) if c.get('categoria')]
    linhas += [('Preço', 'Permuta física') for _, q in ordenados(m['permuta_fisica']) if q.get('unidades') is not None]
    linhas += [('Preço', 'Permuta financeira') for _, q in ordenados(m['permuta_financeira']) if q.get('unidade') is not None]
    return linhas


def candidatas(linhas, cat, sub):
    """Índices das linhas que o rótulo designa: com subcategoria, o par exato; sem ela, a categoria inteira."""
    return [i for i, (c, sc) in enumerate(linhas) if c == cat and (sub is None or sc == sub)]


# Fail-closed: todo campo que NÃO é texto nem Sim/Não é número. Enumerar os numéricos não converge — campo novo
# ficaria sem checagem —; enumerar os não numéricos faz campo novo cair na checagem mais estrita.
# cópia de EVENTOS_ANCORA em backend/rotas/funding.ts
EVENTOS_ANCORA = ('planejamento', 'pre_lancamento', 'lancamento', 'obra', 'pos_obra', 'customizado')
CAMPOS_TEXTO = frozenset((
    'grupo', 'tipologia', 'nome', 'categoria', 'subcategoria', 'cronograma', 'distribuicao', 'unidade', 'linha',
    'aporte', 'modo_retorno', 'base_financiavel', 'tipo', 'tipo_unidade', 'residuo_sem_prazo', 'data_inicio_projeto'))
CAMPOS_SIM_NAO = frozenset((
    'correcao_estoque', 'distribuir_aporte', 'amortizar_com_caixa_disponivel', 'ativo', 'ao_longo_obra',
    'deduzir_corretagem', 'deduzir_imposto', 'considerar_ret', 'tem_pre_lancamento'))


ESTUDO_TEXTO = frozenset(('nome', 'descricao', 'matricula', 'notas', 'nivel_analise', 'origem_terreno',
                          'terreno_manual_nome', 'tipo_empreendimento', 'uf'))
ESTUDO_SIM_NAO = frozenset(('sujeito_ret', 'corretagem_sobre_permuta_fisica'))


def _nao_numero(x):
    return isinstance(x, bool) or not isinstance(x, (int, float)) or x != x or x in (float('inf'), float('-inf'))


def validar(m):
    erros, avisos = [], []
    E, A = erros.append, avisos.append
    e = m.estudo
    # tipo errado (texto onde se espera número) vira erro listado, não traceback; o valor é anulado para
    # o resto da validação seguir e listar os demais problemas
    for lista, itens in m.listas.items():
        for n, v in itens.items():
            for c in list(v):
                x = v[c]
                if c not in CAMPOS_TEXTO and c not in CAMPOS_SIM_NAO and x is not None and _nao_numero(x):
                    E(f'{lista}[{n}].{c}: esperado número, veio "{x}"')
                    v[c] = None
    for ev, d in m.crono.items():
        for c, x in list(d.items()):
            if x is not None and _nao_numero(x):
                E(f'cronograma.{ev}.{c}: esperado número, veio "{x}"')
                d[c] = None
    for c, x in list(m.estudo.items()):
        if c in ESTUDO_TEXTO or c.endswith('_modo') or x is None:
            continue
        if c in ESTUDO_SIM_NAO or c.startswith('considerar_'):
            if not isinstance(x, bool):
                E(f'estudo.{c}: esperado Sim ou Não, veio "{x}"')
                m.estudo[c] = None
        elif _nao_numero(x):
            E(f'estudo.{c}: esperado número, veio "{x}"')
            m.estudo[c] = None
    for c, x in list(m.param.items()):
        if c not in CAMPOS_TEXTO and c not in CAMPOS_SIM_NAO and x is not None and _nao_numero(x):
            E(f'parametros.{c}: esperado número, veio "{x}"')
            m.param[c] = None
    for itens in [{0: m.param}, *m.listas.values()]:
        for n, v in itens.items():
            for c in CAMPOS_SIM_NAO:
                if v.get(c) is not None and not isinstance(v[c], bool):
                    E(f'{c}: esperado Sim ou Não, veio "{v[c]}"')
                    v[c] = None
    for obrig in ('nome', 'tipo_empreendimento', 'nivel_analise'):
        if not e.get(obrig):
            E(f'campo obrigatório ausente: {obrig}')
    if not (isinstance(e.get('terreno_manual_area'), (int, float)) and e['terreno_manual_area'] > 0):
        E('Área do terreno é obrigatória e deve ser > 0')
    if e.get('origem_terreno') == 'nucleo':
        E('origem = nucleo não é suportada pelo importador; use manual')
    for rot, mot in m.nao_enviadas:
        if rot and mot.startswith('vínculo'):
            E(f'"{rot}": {mot}')
        else:
            A(f'não enviado — "{rot}": {mot}')
    # uma só unidade por campo
    for modo_col, rotulos in m.alternativas.items():
        if len(rotulos) > 1:
            E(f'mais de uma unidade preenchida para o mesmo campo ({modo_col}): ' + ' | '.join(rotulos))
    nivel = e.get('nivel_analise')
    if nivel == 'preliminar':
        prods = m['preliminar_produtos']
        if not prods:
            E('Preliminar sem nenhum produto (o catálogo é a única fonte do VGV)')
        soma = 0.0
        for n, p in sorted(prods.items()):
            falta = [c for c in ('nome', 'tipo', 'pct_alv', 'unidades', 'preco_venda_m2') if p.get(c) in (None, '')]
            if falta:
                E(f'Produto {n}: falta {", ".join(falta)}')
            soma += float(p.get('pct_alv') or 0)
        if prods and abs(soma - 100) > 0.01:
            E(f'Σ % da base dos produtos = {soma:.4f}, deve ser 100')
        return erros, avisos
    return _validar_avancado(m, erros, avisos)


def _validar_avancado(m, erros, avisos):
    E, A = erros.append, avisos.append
    e, p = m.estudo, m.param
    lote = e.get('tipo_empreendimento') == 'loteamento'
    for c in ('data_inicio_projeto', 'tem_pre_lancamento'):
        if p.get(c) is None:
            E(f'campo obrigatório ausente: parametros.{c}')
    for ev, campos in (('planejamento', ('inicio_mes', 'duracao_meses')), ('lancamento', ('duracao_meses',)),
                       ('obra', ('inicio_mes', 'duracao_meses')), ('pos_obra', ('duracao_meses',))):
        for c in campos:
            if m.crono.get(ev, {}).get(c) is None:
                E(f'cronograma.{ev}.{c} ausente')
    tem_pre = p.get('tem_pre_lancamento') is True
    if tem_pre and m.crono.get('pre_lancamento', {}).get('duracao_meses') is None:
        E('tem Pré-lançamento = Sim, mas a duração do Pré-lançamento está em branco')
    if not tem_pre and m.crono.get('pre_lancamento', {}).get('duracao_meses') is not None:
        A('duração do Pré-lançamento preenchida com "tem Pré-lançamento = Não": será ignorada')

    def sem_nome(lista, rotulo, campo_nome='nome'):
        """Item com dados mas sem nome seria descartado em silêncio: é erro de preenchimento."""
        nomes = []
        for n, v in sorted(m[lista].items()):
            dados = [k for k, x in v.items() if k != campo_nome and x is not None]
            if v.get(campo_nome) is None and dados:
                E(f'{rotulo} {n}: tem dados ({", ".join(dados[:3])}…) mas o nome está em branco')
            if v.get(campo_nome) is not None:
                nomes.append(str(v[campo_nome]).strip())
        repetidos = sorted({x for x in nomes if nomes.count(x) > 1})
        if repetidos:
            E(f'{rotulo}: nome repetido ({", ".join(repetidos)})')
    sem_nome('tipologias', 'Lote' if lote else 'Tipologia')
    sem_nome('grupos', 'Grupo de venda')
    sem_nome('fases', 'Fase')

    def inteiro(x, minimo):
        return isinstance(x, (int, float)) and float(x).is_integer() and x >= minimo
    for n, f in sorted(m['fases'].items()):
        if f.get('nome'):
            if f.get('inicio_mes') is not None and not inteiro(f['inicio_mes'], 0):
                E(f'Fase {n}: início deve ser um inteiro ≥ 0 (veio {f["inicio_mes"]})')
            if f.get('duracao_meses') is not None and not inteiro(f['duracao_meses'], 1):
                E(f'Fase {n}: duração deve ser um inteiro ≥ 1 (veio {f["duracao_meses"]})')
    sem_nome('funding.divida', 'Dívida')
    sem_nome('funding.equity', 'Equity')
    sem_nome('funding.antecipacao', 'Antecipação de recebíveis')
    sem_nome('cenarios', 'Cenário')
    sem_nome('funding.financiamento_producao', 'Financiamento à produção')

    def sem_chave(lista, rotulo, chaves):
        """Linha com dados mas sem a chave que a identifica seria descartada em silêncio pelo escritor."""
        for n, v in sorted(m[lista].items()):
            dados = [k for k, x in v.items() if k not in chaves and x is not None]
            if all(v.get(c) is None for c in chaves) and dados:
                E(f'{rotulo} {n}: tem dados ({", ".join(dados[:3])}…) mas falta {" / ".join(chaves)}')
    sem_chave('pagamento', 'Condição', ('grupo', 'linha'))
    sem_chave('alocacoes', 'Alocação', ('grupo', 'tipologia'))
    sem_chave('permuta_financeira', 'Permuta financeira', ('unidade',))
    tips = {t['nome']: t for t in m['tipologias'].values() if t.get('nome')}
    if not tips:
        E('nenhuma tipologia/lote informado')
    for n, t in sorted(m['tipologias'].items()):
        if t.get('nome') and (not t.get('quantidade') or not t.get('area_privativa_m2')):
            E(f'{"Lote" if lote else "Tipologia"} {n}: falta área privativa ou unidades')
    grupos = {g['nome']: g for g in m['grupos'].values() if g.get('nome')}
    if not grupos:
        E('nenhum grupo de venda informado')
    if len(grupos) != sum(1 for g in m['grupos'].values() if g.get('nome')):
        E('nomes de grupo de venda repetidos')
    for n, g in sorted(m['grupos'].items()):
        if not g.get('nome'):
            continue
        soma = sum(float(g.get(c) or 0) for c in ('pre_lancamento_pct', 'lancamento_pct', 'obra_pct'))
        if soma > 100.01:
            E(f'Grupo "{g["nome"]}": absorção pré+lançamento+obra = {soma:.2f}% (> 100)')
        janela = g.get('pos_chaves_meses')
        if not (isinstance(janela, (int, float)) and float(janela).is_integer() and 1 <= janela <= 12):
            E(f'Grupo "{g["nome"]}": janela pós-chaves deve ser um inteiro de 1 a 12 (veio {janela!r})')
        if tem_pre and g.get('pre_lancamento_pct') is None:
            A(f'Grupo "{g["nome"]}": há pré-lançamento e a absorção de pré está em branco (vale 0%)')
        if not tem_pre and g.get('pre_lancamento_pct'):
            E(f'Grupo "{g["nome"]}": absorção de pré-lançamento sem pré-lançamento')
        # saldo de chaves (#850) — espelha `erroDestinoSaldoChaves` (frontend/fluxo-pagamento-contrato.ts)
        banco, parc, car = g.get('repasse_banco_pct'), g.get('carteira_parcelas'), g.get('carteira_carencia_meses')
        if banco is not None and not (0 <= banco <= 100):
            E(f'Grupo "{g["nome"]}": repasse bancário deve ficar entre 0% e 100% do saldo de chaves (veio {banco})')
        if parc is not None and not (float(parc).is_integer() and parc >= 1):
            E(f'Grupo "{g["nome"]}": nº de parcelas da carteira própria deve ser um inteiro ≥ 1 (veio {parc})')
        if car is not None and not (float(car).is_integer() and car >= 0):
            E(f'Grupo "{g["nome"]}": carência da carteira própria deve ser um inteiro ≥ 0 (veio {car})')
        if banco is not None and banco < 100 and parc is None:
            E(f'Grupo "{g["nome"]}": repasse bancário abaixo de 100% exige o nº de parcelas da carteira própria')
        if (banco is None or banco >= 100) and (parc is not None or car is not None):
            A(f'Grupo "{g["nome"]}": parcelas/carência da carteira preenchidas com repasse bancário de 100%: '
              'não há carteira própria, e elas são ignoradas')
    por_grupo = defaultdict(float)
    parcelamentos = defaultdict(int)
    for n, q in sorted(m['pagamento'].items()):
        if q.get('grupo') is None and q.get('linha') is None:
            continue
        if q.get('grupo') not in grupos:
            E(f'Condição {n}: grupo "{q.get("grupo")}" não existe na seção de grupos')
            continue
        if q.get('linha') not in ('Entrada', 'Parcelamento'):
            E(f'Condição {n}: Linha deve ser Entrada ou Parcelamento')
        if q.get('pct') in (None, 0):
            E(f'Condição {n}: falta o % do total')
        por_grupo[q['grupo']] += float(q.get('pct') or 0)
        pa = q.get('parcelas')
        if pa is not None and not (float(pa).is_integer() and pa >= 1):
            E(f'Condição {n}: nº de parcelas deve ser um inteiro ≥ 1 (veio {pa})')
        if q.get('linha') == 'Parcelamento':
            parcelamentos[q['grupo']] += 1
            if not q.get('ao_longo_obra') and not q.get('parcelas'):
                E(f'Condição {n}: Parcelamento fora da obra precisa do nº de parcelas')
    for gnome, soma in por_grupo.items():
        if soma > 100.01:
            E(f'Grupo "{gnome}": Σ % do pagamento = {soma:.2f} (> 100)')
        if parcelamentos[gnome] > 4:
            E(f'Grupo "{gnome}": mais de 4 linhas de Parcelamento')
    alocado = defaultdict(int)
    for n, a in sorted(m['alocacoes'].items()):
        if a.get('grupo') is None and a.get('tipologia') is None:
            continue
        if a.get('grupo') not in grupos:
            E(f'Alocação {n}: grupo "{a.get("grupo")}" não existe')
        if a.get('tipologia') not in tips:
            E(f'Alocação {n}: {"lote" if lote else "tipologia"} "{a.get("tipologia")}" não existe')
            continue
        alocado[a['tipologia']] += int(a.get('unidades') or 0)
        if not a.get('unidades') or not a.get('preco_m2'):
            E(f'Alocação {n}: falta unidades ou preço')
    for nome, t in tips.items():
        if alocado[nome] > int(t.get('quantidade') or 0):
            E(f'"{nome}": alocado {alocado[nome]} > catálogo {t.get("quantidade")}')
        elif alocado[nome] < int(t.get('quantidade') or 0):
            A(f'"{nome}": há {int(t["quantidade"]) - alocado[nome]} unidade(s) sem alocar')
    for n, q in sorted(m['permuta_fisica'].items()):
        if q.get('tipologia') is None and q.get('unidades') is None:
            continue
        if not q.get('unidades'):
            E(f'Permuta física {n}: unidades deve ser > 0')
        if q.get('tipologia') not in tips:
            E(f'Permuta física {n}: "{q.get("tipologia")}" não existe')
        elif int(q.get('unidades') or 0) > alocado[q['tipologia']]:
            E(f'Permuta física {n}: {q.get("unidades")} unidades > alocado ({alocado[q["tipologia"]]})')
    # custos
    obrig = {('terreno', 'Preço'), ('obra', 'Construção'), ('diretos', 'Corretagem de vendas')}
    vistas = set()
    for n, c in sorted(m['custos'].items()):
        if c.get('grupo') is None and c.get('categoria') is None:
            continue
        chave = (c.get('grupo'), c.get('categoria'))
        vistas.add(chave)
        un = c.get('unidade')
        if chave not in UNIDADES_POR_CATEGORIA:
            E(f'Custo {n}: categoria "{c.get("categoria")}" inválida no grupo "{c.get("grupo")}"')
            continue
        if un not in UNIDADES_POR_CATEGORIA[chave]:
            E(f'Custo {n}: unidade "{un}" não aceita em {chave[0]}/{chave[1]} (aceitas: {", ".join(UNIDADES_POR_CATEGORIA[chave])})')
            continue
        if c.get('valor_rs') is not None and c.get('valor_pct') is not None:
            E(f'Custo {n}: preencha só uma linha de valor')
        campo = 'valor_pct' if str(un).startswith('pct') else 'valor_rs'
        if c.get(campo) is None:
            E(f'Custo {n}: falta o valor ({campo}) para a unidade {un}')
        elif c[campo] < 0:
            E(f'Custo {n}: {campo} não pode ser negativo (veio {c[campo]})')
        cr0 = str(c.get('cronograma') or '')
        if cr0 and cr0 != 'customizado' and not cr0.startswith('fase:') and cr0 not in EVENTOS:
            E(f'Custo {n}: cronograma "{cr0}" inválido (use um evento, customizado ou fase: <nome>)')
        if c.get('inicio_mes') is not None and not inteiro(c['inicio_mes'], 0):
            E(f'Custo {n}: início deve ser um inteiro ≥ 0 (veio {c["inicio_mes"]})')
        if c.get('duracao_meses') is not None and not inteiro(c['duracao_meses'], 1):
            E(f'Custo {n}: duração deve ser um inteiro ≥ 1 (veio {c["duracao_meses"]})')
        if c.get('cronograma') == 'customizado' and (c.get('inicio_mes') is None or not c.get('duracao_meses')):
            E(f'Custo {n}: cronograma customizado exige início e duração')
    for o in obrig - vistas:
        E(f'custo obrigatório ausente: {o[0]}/{o[1]}')
    fases = {f['nome'] for f in m['fases'].values() if f.get('nome')}
    for n, c in sorted(m['custos'].items()):
        cr = str(c.get('cronograma') or '')
        if cr.startswith('fase:') and cr[5:].strip() not in fases:
            E(f'Custo {n}: fase "{cr[5:].strip()}" não existe nas fases customizadas')
    for n, q in sorted(m['permuta_financeira'].items()):
        un = q.get('unidade')
        if un is None:
            continue
        campo = 'valor_pct' if un == 'pct_vgv' else 'valor_rs'
        if not q.get(campo):
            E(f'Permuta financeira {n}: falta o valor ({campo}) ou ele é 0')
    # funding
    linhas_custo = linhas_de_custo_planejadas(m)
    for n, o in sorted(m['funding.financiamento_producao'].items()):
        for cat, sub in rotulos_da_base(o.get('base_financiavel')):
            achadas = candidatas(linhas_custo, cat, sub)
            rot = f'{cat} / {sub}' if sub else cat
            if len(achadas) == 0:
                E(f'Base financiável: "{rot}" não casa com nenhuma linha de custo')
            elif len(achadas) > 1:
                E(f'Base financiável: "{rot}" é ambíguo ({len(achadas)} linhas); use "Categoria / Subcategoria" para escolher uma só')
    fases_nomes = {f['nome'] for f in m['fases'].values() if f.get('nome')}
    for d in ('funding.divida', 'funding.equity', 'funding.antecipacao'):
        for n, o in sorted(m[d].items()):
            if d != 'funding.antecipacao' and o.get('nome') and not o.get('valor'):
                E(f'{d} {n}: falta o valor')
            ap = str(o.get('aporte') or '')
            if ap.startswith('fase:') and ap[5:].strip() not in fases_nomes:
                E(f'{d} {n}: fase "{ap[5:].strip()}" não existe nas fases customizadas')
            elif ap and not ap.startswith('fase:') and ap not in EVENTOS_ANCORA:
                E(f'{d} {n}: âncora "{ap}" inválida; use um de: {", ".join(EVENTOS_ANCORA)} ou fase:<nome>')
    for n, o in sorted(m['funding.antecipacao'].items()):
        if not o.get('nome'):
            continue
        pc = o.get('pct_cedida')
        if pc is None or not (0 < pc <= 100):
            E(f'funding.antecipacao {n}: carteira cedida deve ficar entre 0% (exclusivo) e 100% (veio {pc!r})')
        if not o.get('aporte'):
            E(f'funding.antecipacao {n}: falta o mês da cessão (evento, fase ou customizado)')
        elif o.get('aporte') == 'customizado' and o.get('mes') is None:
            E(f'funding.antecipacao {n}: mês da cessão customizado exige o mês')
        for c in ('taxa_anual', 'estruturacao_pct', 'administracao_mensal', 'outros_encargos'):
            if o.get(c) is not None and o[c] < 0:
                E(f'funding.antecipacao {n}: {c} não pode ser negativo (veio {o[c]})')
    for n, o in sorted(m['funding.divida'].items()):
        am, ca = o.get('amortizacao_meses'), o.get('carencia_meses')
        if am and ca is not None and am > 0 and ca >= am:
            E(f'funding.divida {n}: carência ({ca}) deve ser menor que a amortização ({am}) — a API recusa')
    por_modo = defaultdict(float)
    for n, o in sorted(m['funding.equity'].items()):
        if o.get('nome'):
            por_modo[str(o.get('modo_retorno') or 'permuta_financeira')] += float(o.get('pct_retorno') or 0)
    for modo, soma in por_modo.items():
        if soma > 100.01:
            E(f'funding.equity: Σ % de retorno em "{modo}" = {soma:.2f} (> 100) — a API recusa a operação que passar de 100')
    return erros, avisos


class EscritaFalhou(Exception):
    """Defeito interno durante a escrita: vai para a mesma limpeza das falhas HTTP, nunca sai direto."""


# ═════════════════════════ HTTP ═════════════════════════
class _SemRedirect(urllib.request.HTTPRedirectHandler):
    """O `Authorization` não pode seguir um redirect para outro host."""
    def redirect_request(self, *a, **k):
        return None


def base_da_api():
    url = os.environ['URBIVERSO_QA_URL'].rstrip('/')
    if not url.startswith('https://'):
        raise SystemExit('URBIVERSO_QA_URL deve começar com https:// (o token não vai em texto claro).')
    return url


class Api:
    def __init__(self, executar):
        self.executar = executar
        self.log = []
        self.abridor = urllib.request.build_opener(_SemRedirect, urllib.request.HTTPSHandler(context=CTX))

    def chamar(self, metodo, caminho, corpo=None):
        if not self.executar:
            self.log.append((metodo, caminho, 'plano', ''))
            return 0, {'id': 0}
        url = base_da_api() + '/api/viabilidade' + caminho
        dados = json.dumps(corpo).encode() if corpo is not None else None
        req = urllib.request.Request(url, data=dados, method=metodo, headers={
            'Authorization': 'Bearer ' + os.environ['URBIVERSO_QA_TOKEN_PCPAL'], 'Content-Type': 'application/json'})
        try:
            r = self.abridor.open(req, timeout=60)
            status, txt = r.status, r.read().decode()
        except urllib.error.HTTPError as ex:
            status, txt = ex.code, ex.read().decode()
        try:
            js = json.loads(txt) if txt else {}
        except ValueError:
            js = {'_texto': txt[:300]}
        self.log.append((metodo, caminho, status, re.sub(r'Bearer\s+\S+', 'Bearer ***', json.dumps(js, ensure_ascii=False))[:240] if status >= 400 else ''))
        return status, js

    def get(self, caminho):
        return self.chamar('GET', caminho)[1]

    def lista(self, caminho):
        js = self.get(caminho)
        return js if isinstance(js, list) else js.get('dados', [])

    def falhas(self):
        return [(m, c, s, r) for m, c, s, r in self.log if isinstance(s, int) and s >= 400]


# ═════════════════════════ escrita ═════════════════════════
def ordenados(d):
    return [(n, v) for n, v in sorted(d.items()) if any(x is not None for x in v.values())]


# falha de transporte (conexão, tempo esgotado, resposta cortada): o pedido pode ter chegado ao servidor
ERROS_REDE = (urllib.error.URLError, OSError, http.client.HTTPException)
CAMPOS_CRIACAO = ('nome', 'tipo_empreendimento', 'nivel_analise', 'origem_terreno', 'uf',
                  'terreno_manual_nome', 'terreno_manual_area', 'notas')


def criar_estudo(api, m):
    """Só o POST: o id precisa existir (e ser impresso) antes de qualquer outra escrita."""
    corpo = {k: m.estudo[k] for k in CAMPOS_CRIACAO if k in m.estudo}
    corpo.setdefault('origem_terreno', 'manual')
    st, r = api.chamar('POST', '/estudos', corpo)
    if st not in (0, 201):
        return None
    return r.get('id') if isinstance(r, dict) else None


def completar_estudo(api, m, eid):
    """Os campos do estudo que não vão no POST; roda já com o estudo criado, dentro da limpeza."""
    resto = {k: v for k, v in m.estudo.items() if k not in CAMPOS_CRIACAO}
    if resto:
        api.chamar('PATCH', f'/estudos/{eid}', resto)


def escrever_preliminar(api, m, eid):
    for n, p in ordenados(m['preliminar_produtos']):
        api.chamar('POST', f'/estudos/{eid}/preliminar/produtos', {
            'nome': p.get('nome'), 'tipo': p.get('tipo'), 'pct_alv': p.get('pct_alv'), 'unidades': p.get('unidades'),
            'preco_venda_m2': p.get('preco_venda_m2'), 'ordem': n - 1})


def escrever_avancado(api, m, eid):
    base = f'/estudos/{eid}/avancado'
    tem_pre = m.param.get('tem_pre_lancamento') is True
    api.chamar('PATCH', f'{base}/parametros', {k.replace('_aa', '_aa'): v for k, v in m.param.items()})
    eventos = {}
    for ev in EVENTOS:
        d = {c: v for c, v in m.crono.get(ev, {}).items() if v is not None}
        # inícios de Pré-lançamento, Lançamento e Pós-obras são travados no fim do evento anterior (o app recalcula)
        if ev in ('pre_lancamento', 'lancamento', 'pos_obra'):
            d.pop('inicio_mes', None)
        if ev == 'pre_lancamento' and not tem_pre:
            continue
        if d:
            eventos[ev] = d
    api.chamar('PATCH', f'{base}/cronograma', {'eventos': eventos})

    fase_id = {}
    for n, f in ordenados(m['fases']):
        if not f.get('nome'):
            continue
        _, r = api.chamar('POST', f'{base}/fases', {'tipo': 'cronograma', 'nome': f['nome'], 'inicio_mes': f.get('inicio_mes', 0),
                                                    'duracao_meses': f.get('duracao_meses', 1), 'ordem': n - 1})
        fase_id[f['nome']] = r.get('id')
    tip_id = {}
    lote = m.estudo.get('tipo_empreendimento') == 'loteamento'
    for n, t in ordenados(m['tipologias']):
        if not t.get('nome'):
            continue
        corpo = {'nome': t['nome'], 'tipo_unidade': 'lote' if lote else t.get('tipo_unidade', 'apartamento'),
                 'area_privativa_m2': t.get('area_privativa_m2'), 'area_privativa_aberta_m2': t.get('area_privativa_aberta_m2', 0),
                 'quantidade': t.get('quantidade', 0), 'preco_m2': t.get('preco_m2'), 'ordem': n - 1}
        if not lote:
            corpo['dormitorios'] = t.get('dormitorios')
            corpo['vagas'] = t.get('vagas')
        _, r = api.chamar('POST', f'{base}/tipologias', corpo)
        tip_id[t['nome']] = r.get('id')

    grupo_id = {}
    for n, g in ordenados(m['grupos']):
        if not g.get('nome'):
            continue
        _, r = api.chamar('POST', f'{base}/fases', {'tipo': 'receita', 'nome': g['nome'], 'ordem': n - 1})
        fid = grupo_id[g['nome']] = r.get('id')
        pre, lanc, obra = (float(g.get(c) or 0) for c in ('pre_lancamento_pct', 'lancamento_pct', 'obra_pct'))
        blocos = ([{'evento': 'pre_lancamento', 'pct': pre}] if tem_pre else []) + [
            {'evento': 'lancamento', 'pct': lanc}, {'evento': 'obra', 'pct': obra},
            {'evento': 'pos_obra', 'pct': round(100 - pre - lanc - obra, 6)}]      # pós-chaves já DERIVADO
        absorcao = {'modo': 'distribuido', 'correcao_estoque': bool(g.get('correcao_estoque', False)),
                    'pos_chaves_meses': int(g.get('pos_chaves_meses') or 12), 'blocos': blocos}
        entradas, parcelas = [], []
        for _, q in ordenados(m['pagamento']):
            if q.get('grupo') != g['nome']:
                continue
            if q.get('linha') == 'Entrada':
                entradas.append({'pct': q.get('pct', 0), 'parcelas': q.get('parcelas', 1), 'descontoPct': q.get('desconto_pct', 0)})
            else:
                parcelas.append({'periodicidade': 'mensal', 'parcelas': q.get('parcelas', 0), 'ao_longo_obra': bool(q.get('ao_longo_obra', False)),
                                 'juros': True, 'pct': q.get('pct', 0), 'sinalPct': q.get('sinal_pct', 0)})
        repasse = {'apos_entrega_meses': 1}
        banco = g.get('repasse_banco_pct')
        if banco is not None and banco < 100:          # carteira própria (#850); vazio ou 100% = tudo no banco, como antes
            repasse.update(pct_banco=banco, carteira_parcelas=int(g['carteira_parcelas']),
                           carteira_carencia_meses=int(g.get('carteira_carencia_meses') or 0))
        fluxo = {'entrada': entradas, 'parcelas': parcelas, 'repasse': repasse,
                 'residuoAteMarco': g.get('residuo_sem_prazo', 'imediato')}
        api.chamar('PATCH', f'{base}/fases/{fid}', {'absorcao': absorcao, 'fluxo_pagamento': fluxo})
    for n, a in ordenados(m['alocacoes']):
        if a.get('grupo') is None:
            continue
        api.chamar('POST', f'{base}/fases/{grupo_id.get(a["grupo"])}/alocacoes',
                   {'tipologia_id': tip_id.get(a.get('tipologia')), 'unidades': a.get('unidades', 0), 'preco_m2': a.get('preco_m2')})

    curvas = {c['nome']: c['id'] for c in api.lista('/avancado/curvas')} if api.executar else {}
    custo_ids = []                      # id de cada linha de custo criada, na ordem de linhas_de_custo_planejadas
    linhas_custo = linhas_de_custo_planejadas(m)

    def comum(c):
        d = {}
        dist = c.get('distribuicao')
        if dist in curvas:
            d.update(curva_id=curvas[dist], distribuicao_modo='fixo')
        elif dist in ('unit_delivery', 'sales_revenue'):
            d['distribuicao_modo'] = dist
        elif dist in ('Linear', 'fixo'):
            d['distribuicao_modo'] = 'fixo'
        elif dist:
            if api.executar:                  # curva que a instância não tem: nunca vira Linear em silêncio
                raise EscritaFalhou(f'distribuição "{dist}" não existe no catálogo de curvas da instância')
            d['distribuicao_modo'] = 'fixo'
        cr = c.get('cronograma')
        if cr and str(cr).startswith('fase:'):
            d['fase_ancora_id'] = fase_id.get(str(cr)[5:].strip())
        elif cr:
            d['cronograma_evento'] = cr
            if cr == 'customizado':
                d.update(inicio_mes=c.get('inicio_mes', 0), duracao_meses=c.get('duracao_meses', 1))
        return d

    for n, c in ordenados(m['custos']):
        if c.get('categoria') is None:
            continue
        un = c['unidade']
        valor = c.get('valor_pct') if str(un).startswith('pct') else c.get('valor_rs')
        corpo = {'grupo': c['grupo'], 'categoria': c['categoria'], 'subcategoria': c.get('subcategoria'),
                 'orcamento_unidade': un, 'orcamento_valor': valor, 'ordem': n, **comum(c)}
        _, r = api.chamar('POST', f'{base}/custos', corpo)
        custo_ids.append(r.get('id'))
    for n, q in ordenados(m['permuta_fisica']):
        if q.get('unidades') is None:
            continue
        _, r = api.chamar('POST', f'{base}/custos', {'grupo': 'terreno', 'categoria': 'Preço', 'subcategoria': 'Permuta física',
                   'permuta_tipologia_id': tip_id.get(q.get('tipologia')), 'permuta_quantidade': q['unidades'],
                   'distribuicao_modo': 'unit_delivery', 'ordem': 100 + n})
        custo_ids.append(r.get('id'))
    for n, q in ordenados(m['permuta_financeira']):
        if q.get('unidade') is None:
            continue
        valor = q.get('valor_pct') if q['unidade'] == 'pct_vgv' else q.get('valor_rs')
        _, r = api.chamar('POST', f'{base}/custos', {
            'grupo': 'terreno', 'categoria': 'Preço', 'subcategoria': 'Permuta financeira', 'orcamento_unidade': q['unidade'],
            'orcamento_valor': valor, 'distribuicao_modo': 'sales_revenue', 'ordem': 110 + n,
            'permuta_financeira_deduzir_imposto': bool(q.get('deduzir_imposto', False)),
            'permuta_financeira_deduzir_corretagem': bool(q.get('deduzir_corretagem', False))})
        custo_ids.append(r.get('id'))

    def anc(o):
        a = o.get('aporte')
        if a and str(a).startswith('fase:'):
            return {'fase_ancora_id': fase_id.get(str(a)[5:].strip())}
        return {'cronograma_evento': a} if a else {}

    for n, o in ordenados(m['funding.financiamento_producao']):
        if not o.get('nome'):
            continue
        d = {'tipo': 'financiamento_producao', 'nome': o['nome'], 'ordem': 0}
        for c in ('ativo', 'valor', 'taxa_anual', 'exposicao_minima', 'percentual_financiavel', 'amortizar_com_caixa_disponivel'):
            if o.get(c) is not None:
                d[c] = o[c]
        base_fin = o.get('base_financiavel')
        if base_fin:
            ids = []
            for cat, sub in rotulos_da_base(base_fin):
                achadas = candidatas(linhas_custo, cat, sub)
                if len(achadas) == 1 and achadas[0] < len(custo_ids) and (custo_ids[achadas[0]] or not api.executar):
                    ids.append(custo_ids[achadas[0]])
                else:                          # a validação já barrou isto; chegar aqui é defeito, não silêncio
                    raise EscritaFalhou(f'base financiável "{cat}": não resolveu uma única linha de custo criada')
            d['custo_linha_ids'] = ids
        api.chamar('POST', f'{base}/funding', d)
    for n, o in ordenados(m['funding.divida']):
        if not o.get('nome'):
            continue
        d = {'tipo': 'divida', 'nome': o['nome'], 'ordem': n, **anc(o)}
        for c, col in (('valor', 'valor'), ('mes', 'inicio_mes'), ('distribuir_aporte', 'distribuir_aporte'), ('aporte_meses', 'aporte_meses'),
                       ('taxa_anual', 'taxa_anual'), ('amortizacao_meses', 'periodo_amortizacao_meses'),
                       ('carencia_meses', 'periodo_carencia_meses'), ('estruturacao_pct', 'taxa_estruturacao_pct'),
                       ('administracao_mensal', 'taxa_administracao_mensal'), ('outros_encargos', 'outros_encargos_iniciais')):
            if o.get(c) is not None:
                d[col] = o[c]
        api.chamar('POST', f'{base}/funding', d)
    for n, o in ordenados(m['funding.equity']):
        if not o.get('nome'):
            continue
        d = {'tipo': 'equity', 'nome': o['nome'], 'ordem': n, **anc(o)}
        for c, col in (('valor', 'valor'), ('mes', 'inicio_mes'), ('modo_retorno', 'modo_retorno'), ('pct_retorno', 'pct_retorno')):
            if o.get(c) is not None:
                d[col] = o[c]
        api.chamar('POST', f'{base}/funding', d)
    for n, o in ordenados(m['funding.antecipacao']):
        if not o.get('nome'):
            continue
        d = {'tipo': 'antecipacao_recebiveis', 'nome': o['nome'], 'ordem': n, 'valor': 0, **anc(o)}
        for c, col in (('mes', 'inicio_mes'), ('pct_cedida', 'pct_carteira_cedida'), ('taxa_anual', 'taxa_anual'),
                       ('estruturacao_pct', 'taxa_estruturacao_pct'), ('administracao_mensal', 'taxa_administracao_mensal'),
                       ('outros_encargos', 'outros_encargos_iniciais')):
            if o.get(c) is not None:
                d[col] = o[c]
        api.chamar('POST', f'{base}/funding', d)
    for n, c in ordenados(m['cenarios']):
        if c.get('nome'):
            api.chamar('POST', f'{base}/cenarios', {'nome': c['nome'], 'preco_venda_pct': c.get('preco_venda_pct', 0),
                                                    'custo_obra_pct': c.get('custo_obra_pct', 0), 'ordem': n - 1})


def canonico(eid):
    env = dict(os.environ, URBI_BASE=base_da_api(), URBI_TOKEN=os.environ['URBIVERSO_QA_TOKEN_PCPAL'])
    try:
        r = subprocess.run(['node', '--import', 'tsx/esm', 'scripts/formularios/aplicar-pagamento-canonico.ts', str(eid), '--forcar'],
                           cwd=RAIZ, env=env, capture_output=True, text=True, timeout=180)
    except (OSError, subprocess.TimeoutExpired) as ex:
        return 127, f'não consegui rodar o node: {ex}'
    return r.returncode, (r.stdout + r.stderr).strip()


# ═════════════════════════ releitura ═════════════════════════
def _num(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def igual(enviado, gravado, tol=0.0051):
    if isinstance(enviado, bool) or isinstance(gravado, bool):
        return bool(enviado) == bool(gravado)
    a, b = _num(enviado), _num(gravado)
    if a is not None and b is not None and not isinstance(enviado, str):
        return abs(a - b) <= tol
    return str(enviado).strip() == str(gravado).strip()


def reler(api, m, eid):
    """Compara o que foi gravado com o que a planilha pediu. Ausência de resposta é DIVERGÊNCIA, nunca silêncio."""
    div = []
    est = api.get(f'/estudos/{eid}')
    if not est or 'id' not in est:
        return [f'GET /estudos/{eid} não devolveu o estudo: nada foi conferido']
    enviados = {k: v for k, v in m.estudo.items() if k not in ('nivel_analise', 'tipo_empreendimento')}
    for col, v in enviados.items():
        if col not in est:
            div.append(f'estudos.{col}: a coluna não existe na resposta do app')
        elif not igual(v, est[col], tol=0.0051 if col != 'pct_alv' else 0.00006):
            div.append(f'estudos.{col}: enviado {v!r}, gravado {est[col]!r}')
    if m.estudo.get('nivel_analise') == 'preliminar':
        gravados = sorted(api.lista(f'/estudos/{eid}/preliminar/produtos'), key=lambda p: p.get('ordem', 0))
        pedidos = [v for _, v in ordenados(m['preliminar_produtos'])]
        if len(gravados) != len(pedidos):
            div.append(f'produtos: pedidos {len(pedidos)}, gravados {len(gravados)}')
        for p, g in zip(pedidos, gravados):
            for c in ('nome', 'tipo', 'unidades', 'preco_venda_m2'):
                if not igual(p.get(c), g.get(c)):
                    div.append(f'produto "{p.get("nome")}".{c}: enviado {p.get(c)!r}, gravado {g.get(c)!r}')
            if not igual(p.get('pct_alv'), g.get('pct_alv'), tol=0.00006):
                div.append(f'produto "{p.get("nome")}".pct_alv: enviado {p.get("pct_alv")!r}, gravado {g.get("pct_alv")!r}')
        return div
    par = api.get(f'/estudos/{eid}/avancado/parametros')
    for c, v in m.param.items():
        if c not in par:
            div.append(f'parametros.{c}: a API não devolveu o campo (enviado {v!r})')
        elif not igual(v, par[c]):
            div.append(f'parametros.{c}: enviado {v!r}, gravado {par[c]!r}')
    cr = {x['evento']: x for x in api.lista(f'/estudos/{eid}/avancado/cronograma')}
    tem_pre = m.param.get('tem_pre_lancamento') is True
    for ev, campos in m.crono.items():
        for c, v in campos.items():
            if v is None or (ev == 'pre_lancamento' and not tem_pre):
                continue
            if ev not in cr or c not in cr[ev]:
                div.append(f'cronograma.{ev}.{c}: a API não devolveu o campo (enviado {v!r})')
            elif not (c == 'inicio_mes' and cr[ev].get('travado_inicio')) and not igual(v, cr[ev][c]):
                div.append(f'cronograma.{ev}.{c}: enviado {v!r}, gravado {cr[ev][c]!r}')
    tips = api.lista(f'/estudos/{eid}/avancado/tipologias')
    pedidas = [v for _, v in ordenados(m['tipologias']) if v.get('nome')]
    if len(tips) != len(pedidas):
        div.append(f'tipologias: pedidas {len(pedidas)}, gravadas {len(tips)}')
    por_nome = {t['nome']: t for t in tips}
    for t in pedidas:
        g = por_nome.get(t['nome'])
        if not g:
            div.append(f'tipologia "{t["nome"]}" não foi gravada')
            continue
        for c in ('area_privativa_m2', 'area_privativa_aberta_m2', 'quantidade', 'preco_m2'):
            if t.get(c) is not None and not igual(t[c], g.get(c)):
                div.append(f'tipologia "{t["nome"]}".{c}: enviado {t[c]!r}, gravado {g.get(c)!r}')
    receitas = api.lista(f'/estudos/{eid}/avancado/receitas')
    pedidos_g = {v['nome']: v for _, v in ordenados(m['grupos']) if v.get('nome')}
    if len(receitas) != len(pedidos_g):
        div.append(f'grupos: pedidos {len(pedidos_g)}, gravados {len(receitas)}')
    for r in receitas:
        nome = r.get('nome')
        g = pedidos_g.get(nome)
        if g is None:
            div.append(f'grupo "{nome}" existe no estudo e a planilha não o pediu')
            continue
        fp = r.get('fluxo_pagamento') or {}
        comp = fp.get('componentes')
        if not (isinstance(comp, list) and comp):
            div.append(f'grupo "{nome}": fluxo de pagamento SEM componentes (o motor calcula sem juros)')
        banco_ped = g.get('repasse_banco_pct')
        if banco_ped is not None and banco_ped < 100:
            rep = fp.get('repasse') if isinstance(fp.get('repasse'), dict) else {}
            for c, ped in (('pct_banco', banco_ped), ('carteira_parcelas', g.get('carteira_parcelas')),
                           ('carteira_carencia_meses', g.get('carteira_carencia_meses') or 0)):
                if not igual(ped, rep.get(c)):
                    div.append(f'grupo "{nome}": repasse.{c} pedido {ped!r}, gravado {rep.get(c)!r}')
            cart = [c for c in (comp or []) if isinstance(c, dict) and c.get('tipo') == 'carteira_pos_marco']
            if isinstance(comp, list) and comp and not cart:
                div.append(f'grupo "{nome}": carteira própria pedida, mas os componentes não têm carteira_pos_marco')
        else:
            ent = [(float(q.get('pct') or 0), 0.0) for _, q in ordenados(m['pagamento']) if q.get('grupo') == nome and q.get('linha') == 'Entrada']
            par_ = [(float(q.get('pct') or 0), float(q.get('sinal_pct') or 0)) for _, q in ordenados(m['pagamento']) if q.get('grupo') == nome and q.get('linha') == 'Parcelamento']
            repasse = round(100 - sum(x[0] for x in ent + par_), 6)
            banco = g.get('repasse_banco_pct')
            banco = 100.0 if banco is None or banco >= 100 else float(banco)
            saldo = [repasse * banco / 100, repasse * (100 - banco) / 100]     # banco (concentrado) e carteira própria
            esperado = sorted(round(x[0], 2) for x in ent + par_ + [(v, 0.0) for v in saldo if v > 0.005])
            gravado = sorted(round(float(c.get('participacaoPct') or 0), 2) for c in comp)
            if esperado != gravado:
                div.append(f'grupo "{nome}": participações do pagamento pedidas {esperado}, gravadas {gravado}')
            sinal_esp = round(sum(a * b / 100 for a, b in par_), 4)
            sinal_grv = round(sum(float(c.get('participacaoPct') or 0) * float(c.get('sinalPct') or 0) / 100 for c in comp), 4)
            if abs(sinal_esp - sinal_grv) > 0.01:
                div.append(f'grupo "{nome}": sinal ponderado pedido {sinal_esp}, gravado {sinal_grv}')
        blocos = (r.get('absorcao') or {}).get('blocos', [])
        pos = [b for b in blocos if b.get('evento') == 'pos_obra']
        esperado_pos = 100 - sum(float(g.get(c) or 0) for c in ('pre_lancamento_pct', 'lancamento_pct', 'obra_pct'))
        if not pos:
            div.append(f'grupo "{nome}": sem bloco de pós-chaves na absorção')
        elif not igual(esperado_pos, pos[0].get('pct'), tol=0.01):
            div.append(f'grupo "{nome}": pós-chaves gravado {pos[0].get("pct")!r}, derivado {esperado_pos!r}')
        for evento, campo in (('lancamento', 'lancamento_pct'), ('obra', 'obra_pct')):
            b = next((x for x in blocos if x.get('evento') == evento), None)
            if b is None or not igual(g.get(campo) or 0, b.get('pct'), tol=0.01):
                div.append(f'grupo "{nome}": absorção {evento} pedida {g.get(campo)!r}, gravada {None if b is None else b.get("pct")!r}')
        ab = r.get('absorcao') or {}
        if not igual(g.get('pos_chaves_meses', 12), ab.get('pos_chaves_meses', 12)):
            div.append(f'grupo "{nome}": janela pós-chaves pedida {g.get("pos_chaves_meses")!r}, gravada {ab.get("pos_chaves_meses")!r}')
        if bool(g.get('correcao_estoque', False)) != bool(ab.get('correcao_estoque', False)):
            div.append(f'grupo "{nome}": correção de estoque pedida {g.get("correcao_estoque")!r}, gravada {ab.get("correcao_estoque")!r}')
        if not igual(g.get('residuo_sem_prazo', 'imediato'), fp.get('residuoAteMarco', 'imediato')):
            div.append(f'grupo "{nome}": resíduo sem prazo pedido {g.get("residuo_sem_prazo")!r}, gravado {fp.get("residuoAteMarco")!r}')
        gravado_aloc = defaultdict(int)
        for t in r.get('tipologias', []):
            gravado_aloc[t.get('nome') or t.get('tipologia_id')] += int(t.get('quantidade') or 0)
        pedido_aloc = defaultdict(int)
        for _, a in ordenados(m['alocacoes']):
            if a.get('grupo') == nome:
                pedido_aloc[a.get('tipologia')] += int(a.get('unidades') or 0)
        if sum(gravado_aloc.values()) != sum(pedido_aloc.values()):
            div.append(f'grupo "{nome}": unidades alocadas pedidas {sum(pedido_aloc.values())}, gravadas {sum(gravado_aloc.values())}')
        elif all(isinstance(k, str) for k in gravado_aloc) and dict(gravado_aloc) != dict(pedido_aloc):
            div.append(f'grupo "{nome}": alocação por tipologia pedida {dict(pedido_aloc)}, gravada {dict(gravado_aloc)}')
    custos = api.lista(f'/estudos/{eid}/avancado/custos')
    esperados = []
    for _, c in ordenados(m['custos']):
        if c.get('categoria') is None:
            continue
        un = c['unidade']
        esperados.append((c['grupo'], c['categoria'], c.get('subcategoria') or None, un,
                          c.get('valor_pct') if str(un).startswith('pct') else c.get('valor_rs')))
    for _, q in ordenados(m['permuta_fisica']):
        if q.get('unidades') is not None:
            esperados.append(('terreno', 'Preço', 'Permuta física', 'rs', None))
    for _, q in ordenados(m['permuta_financeira']):
        if q.get('unidade') is not None:
            esperados.append(('terreno', 'Preço', 'Permuta financeira', q['unidade'],
                              q.get('valor_pct') if q['unidade'] == 'pct_vgv' else q.get('valor_rs')))
    gravados = [(c['grupo'], c['categoria'], c.get('subcategoria') or None, c['orcamento_unidade'], c.get('orcamento_valor')) for c in custos]
    if len(gravados) != len(esperados):
        div.append(f'custos: pedidos {len(esperados)}, gravados {len(gravados)}')
    restantes = list(gravados)
    for e in esperados:
        achou = next((g for g in restantes if g[:4] == e[:4] and (e[4] is None or igual(e[4], g[4]))), None)
        if achou is None:
            div.append(f'custo {e[0]}/{e[1]}{"/" + e[2] if e[2] else ""} ({e[3]}={e[4]!r}) não foi gravado como pedido')
        else:
            restantes.remove(achou)
    ops = api.lista(f'/estudos/{eid}/avancado/funding')
    pedidas_f = [(t, o) for t, lista in (('financiamento_producao', 'funding.financiamento_producao'), ('divida', 'funding.divida'),
                                         ('equity', 'funding.equity'), ('antecipacao_recebiveis', 'funding.antecipacao'))
                 for _, o in ordenados(m[lista]) if o.get('nome')]
    if len(ops) != len(pedidas_f):
        div.append(f'funding: pedidas {len(pedidas_f)}, gravadas {len(ops)}')
    for tipo, o in pedidas_f:
        g = next((x for x in ops if x.get('tipo') == tipo and x.get('nome') == o['nome']), None)
        if g is None:
            div.append(f'funding {tipo} "{o["nome"]}" não foi gravado')
        elif o.get('valor') is not None and not igual(o['valor'], g.get('valor')):
            div.append(f'funding {tipo} "{o["nome"]}".valor: enviado {o["valor"]!r}, gravado {g.get("valor")!r}')
        if tipo == 'antecipacao_recebiveis' and g is not None:
            for c, col in (('pct_cedida', 'pct_carteira_cedida'), ('taxa_anual', 'taxa_anual')):
                if o.get(c) is not None and not igual(o[c], g.get(col)):
                    div.append(f'funding antecipação "{o["nome"]}".{col}: enviado {o[c]!r}, gravado {g.get(col)!r}')
        if tipo == 'financiamento_producao' and o.get('base_financiavel') and g is not None:
            pedidos_ids = len(rotulos_da_base(o['base_financiavel']))
            if len(g.get('custo_linha_ids') or []) != pedidos_ids:
                div.append(f'financiamento à produção: base financiável pedida com {pedidos_ids} linha(s), gravada com {len(g.get("custo_linha_ids") or [])}')
    cen = api.lista(f'/estudos/{eid}/avancado/cenarios')
    ped = [c for _, c in ordenados(m['cenarios']) if c.get('nome')]
    if len(cen) != len(ped):
        div.append(f'cenários: pedidos {len(ped)}, gravados {len(cen)}')
    return div


def conferir(eid):
    env = dict(os.environ, URBI_BASE=base_da_api(), URBI_TOKEN=os.environ['URBIVERSO_QA_TOKEN_PCPAL'])
    try:
        r = subprocess.run(['node', '--import', 'tsx/esm', 'scripts/conferir-estudo.ts', str(eid)], cwd=RAIZ, env=env,
                           capture_output=True, text=True, timeout=300)
    except (OSError, subprocess.TimeoutExpired) as ex:
        return [], [f'não consegui rodar o motor do app: {ex}'], [], 127
    saida = r.stdout + r.stderr
    resumo, reais, conhecidos = [], [], []
    for ln in saida.splitlines():
        s = ln.strip()
        if re.match(r'^(VGV|Receita líquida|Custo (direto|indireto|total)|Resultado|Margem|ROI|TIR|VPL|Payback|Exposição|Juros de clientes|Carteira|Repasse)', s):
            resumo.append(s)
        if s.startswith('FALHA') or s.startswith('[erro]') or s.startswith('[alerta]'):
            rotulo = next((c for rx, c in CONHECIDOS if rx.search(s)), None)
            (conhecidos if rotulo else reais).append(f'{s[:200]}' + (f'  → {rotulo}' if rotulo else ''))
    reais = [x for x in reais if not x.startswith('[alerta]') or 'CAIXA_ACUMULADO' not in x]
    return resumo, reais, conhecidos, r.returncode


# ═════════════════════════ main ═════════════════════════
def main(argv):
    ap = argparse.ArgumentParser(description='Cria um estudo a partir de um formulário simplificado e confere o resultado.')
    ap.add_argument('formulario', help='o .xlsx preenchido')
    ap.add_argument('--nome-estudo', help='sobrescreve o nome do estudo')
    ap.add_argument('--executar', action='store_true', help='escreve na instância (sem isto só valida e mostra o plano)')
    ap.add_argument('--relatorio', help='grava o relatório em JSON')
    ap.add_argument('--manter-em-falha', action='store_true', help='não remove o estudo quando a API recusa alguma escrita (diagnóstico)')
    a = ap.parse_args(argv[1:])
    executar, rel, arquivo, manter = a.executar, a.relatorio, a.formulario, a.manter_em_falha
    m = montar(ler(arquivo))
    if a.nome_estudo:
        m.estudo['nome'] = a.nome_estudo
    erros, avisos = validar(m)
    relatorio = {'arquivo': arquivo, 'erros_de_preenchimento': erros, 'avisos': avisos,
                 'ignoradas': m.ignoradas, 'estudo_id': None}
    print(f'Formulário: {arquivo}')
    for av in avisos:
        print('  aviso :', av)
    for rot, k in m.ignoradas:
        print(f'  ignorado (sem tratamento): "{rot}" [{k}]')
    for er in erros:
        print('  ERRO  :', er)
    if erros:
        print(f'\n{len(erros)} erro(s) de preenchimento — nada foi escrito.')
        if rel:
            Path(rel).write_text(json.dumps(relatorio, ensure_ascii=False, indent=2))
        return 1
    if executar and rel:
        # o destino do relatório é testado ANTES de criar o estudo: não adianta descobrir depois que não dá para gravar
        try:
            Path(rel).write_text(json.dumps(relatorio, ensure_ascii=False, indent=2))
        except OSError as ex:
            print(f'ERRO: não consigo gravar o relatório em "{rel}" ({ex}) — nada foi escrito na instância.')
            return 1
    api = Api(executar)
    try:
        eid = criar_estudo(api, m)
    except ERROS_REDE as ex:
        # a resposta do POST pode ter se perdido com o estudo já criado: sem id, só dá para conferir pela lista
        print(f'ERRO: falha de rede ao criar o estudo ({type(ex).__name__}: {ex}). Se o pedido chegou ao servidor, '
              'um estudo pode ter sido criado sem id: compare a lista de estudos do app com a de antes de rodar.')
        return 1
    if not eid and eid != 0:
        print('ERRO: o POST /estudos falhou ou voltou sem id (o estudo pode ter sido criado; confira a lista de '
              'estudos do app contra a de antes de rodar):', api.falhas())
        return 1
    if executar:
        # o id sai AQUI, antes de qualquer outra escrita: se a execução cair depois, ele já está no terminal
        # e no arquivo de --relatorio, e dá para achar (e apagar) o estudo parcial
        print(f'estudo criado: id {eid} (nome "{m.estudo.get("nome")}")', flush=True)
        relatorio['estudo_id'] = eid
        if rel:
            try:
                Path(rel).write_text(json.dumps(relatorio, ensure_ascii=False, indent=2))
            except OSError as ex:
                print(f'aviso: não consegui gravar o id em "{rel}" ({ex}); o id está só neste terminal.')

    def desfazer(motivo):
        # a validação prévia não conhece toda regra da API; se uma escrita for recusada depois do estudo existir,
        # o estudo parcial é removido em vez de ficar como se estivesse certo
        print(f'ERRO: {motivo} — estudo parcial:', api.falhas())
        try:
            st, _ = api.chamar('DELETE', f'/estudos/{eid}')
        except ERROS_REDE as ex:
            st = f'{type(ex).__name__}: {ex}'
        print(f'estudo {eid} removido (HTTP {st}); nada ficou gravado pela metade.' if st in (200, 204)
              else f'ATENÇÃO: não consegui remover o estudo {eid} ({st}); apague-o à mão.')
        return 1
    try:
        completar_estudo(api, m, eid)
        (escrever_preliminar if m.estudo['nivel_analise'] == 'preliminar' else escrever_avancado)(api, m, eid)
    except EscritaFalhou as ex:
        if executar and not manter:
            return desfazer(str(ex))
        raise SystemExit(f'{ex}' + (f' (estudo {eid} mantido)' if executar else ''))
    except ERROS_REDE as ex:
        motivo = f'falha de rede durante a escrita ({type(ex).__name__}: {ex})'
        if executar and not manter:
            return desfazer(motivo)
        raise SystemExit(f'{motivo}' + (f' (estudo {eid} mantido)' if executar else ''))
    if executar and api.falhas() and not manter:
        return desfazer('a API recusou escritas')
    for mt, c, st, r_ in api.log:
        if st == 'aviso':
            print('  aviso :', c)
    print(f'\nestudo_id = {eid} · {len(api.log)} chamadas · {len(api.falhas())} falha(s) HTTP')
    if not executar:
        print('\n(plano — use --executar para escrever)')
        return 0
    pendencias = []
    if m.estudo['nivel_analise'] == 'avancado':
        rc, saida = canonico(eid)
        print('\npagamento canônico:', 'ok' if rc == 0 else 'FALHOU')
        for ln in saida.splitlines():
            print('  ', ln)
        if rc != 0:
            if not manter:
                return desfazer('a conversão do pagamento para o formato canônico falhou')
            pendencias.append('a conversão do pagamento para o formato canônico falhou')
    div = reler(api, m, eid)
    falhas = api.falhas()                       # depois da releitura: GET com erro também conta
    print(f'\nreleitura: {len(div)} divergência(s) · {len(falhas)} falha(s) HTTP no total')
    for mt, c, st, r_ in falhas:
        print(f'  XX {mt} {c} {st} {r_}')
    for d in div:
        print('  DIVERGE:', d)
    resumo, reais, conhecidos, rc_motor = conferir(eid)
    print('\nmotor do app:')
    for ln in resumo:
        print('  ', ln)
    for x in conhecidos:
        print('  conhecido:', x)
    for x in reais:
        print('  PROBLEMA :', x)
    if rc_motor != 0:
        pendencias.append(f'o motor do app (conferir-estudo) saiu com código {rc_motor}')
    if not resumo:
        pendencias.append('o motor do app não devolveu nenhum número reconhecível: o estudo NÃO foi conferido')
    for pd in pendencias:
        print('  PENDÊNCIA:', pd)
    relatorio.update(estudo_id=eid, falhas_http=[list(f) for f in falhas], divergencias=div, motor=resumo,
                     problemas=reais, conhecidos=conhecidos, pendencias=pendencias)
    if rel:
        Path(rel).write_text(json.dumps(relatorio, ensure_ascii=False, indent=2))
    ok = not falhas and not div and not reais and not pendencias
    print('\nRESULTADO:', 'OK' if ok else 'COM PENDÊNCIAS')
    return 0 if ok else 2


if __name__ == '__main__':
    sys.exit(main(sys.argv))
